#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
KTDGKQHT_05 — điền tệp mẫu điểm danh thành fixture "3 dòng hợp lệ + 2 dòng cố ý sai".

🔴 KHÔNG tự tạo tệp từ đầu. Tệp mẫu BẮT BUỘC tải từ chính nút "Tải mẫu điểm danh"
   của hệ thống trong lượt đo, vì nó mang sheet metadata `_HTPLDN_META` chứa
   `lich_hoc_id` định danh buổi học. Tệp tự dựng sẽ vướng ERR-KQ-09 → FAIL oan.
   (Đặc tả: srs-fr-03-dao-tao.md:583 metadata tệp · :590 bắt buộc khớp buổi đang chọn)

Script này CHỈ đọc tệp mẫu thật rồi ghi thêm dữ liệu; nó không sinh cấu trúc mới.

CÁCH DÙNG
---------
  # 1) Xem cấu trúc tệp mẫu vừa tải về (không ghi gì)
  python3 KTDGKQHT_05-dien-file-mau.py ~/Downloads/mau-diem-danh-<...>.xlsx --chi-doc

  # 2) Điền fixture (mặc định ghi ra ../files/fixture-KTDGKQHT_05-3hople-2loi-<mốc giờ>.xlsx)
  python3 KTDGKQHT_05-dien-file-mau.py ~/Downloads/mau-diem-danh-<...>.xlsx

  # 3) Chỉ định nơi ghi + học viên khóa khác khác mặc định
  python3 KTDGKQHT_05-dien-file-mau.py <tệp mẫu> -o <tệp ra>.xlsx \
      --hv-khoa-khac <uuid> --ten-hv-khoa-khac "<họ tên>" --email-hv-khoa-khac "<email>"

BỐ CỤC FIXTURE SINH RA (giữ nguyên như lượt đo 02:20 ngày 07/08/2026 để so được số dòng)
----------------------------------------------------------------------------------------
  dòng 2 = học viên 1 của khóa · "Có mặt"            → hợp lệ
  dòng 3 = học viên 2 của khóa · "Vắng có phép"      → hợp lệ
  dòng 4 = học viên 3 của khóa · "Vắng không phép"   → hợp lệ
  dòng 5 = học viên KHÓA KHÁC  · "Có mặt"            → lỗi ERR-KQ-03 (srs-fr-03-dao-tao.md:592, :637)
  dòng 6 = học viên 4 của khóa · "XYZ"               → lỗi ERR-KQ-04 (srs-fr-03-dao-tao.md:638)
  ⇒ Xem trước phải ra: Tổng 5 · Hợp lệ 3 · Bỏ qua 0 · Lỗi 2

Yêu cầu: python3 -m pip install openpyxl
"""

import argparse
import datetime as _dt
import sys
import zipfile
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("Thiếu thư viện openpyxl. Cài bằng: python3 -m pip install openpyxl")


# ---------------------------------------------------------------- hằng số đặc tả
SHEET_META = "_HTPLDN_META"
COT_CHUAN = [
    "hoc_vien_id",
    "Họ tên",
    "Email",
    "Đơn vị",
    "Trạng thái điểm danh",
    "Ghi chú",
]  # srs-fr-03-dao-tao.md:581 — thứ tự cố định

TRANG_THAI_HOP_LE = ["Có mặt", "Vắng có phép", "Vắng không phép"]  # :549, :581, :1919
TRANG_THAI_SAI = "XYZ"  # ép ERR-KQ-04

# Học viên thuộc khóa KHÁC — đã dùng ở lượt đo 02:20 ngày 07/08/2026
# (khóa DDD-KH-011, trạng thái DA_DUYET). 🔴 PHẢI kiểm lại còn tồn tại và vẫn
# thuộc khóa khác trước khi dùng; nếu không còn thì truyền --hv-khoa-khac id mới.
HV_KHOA_KHAC_MAC_DINH = {
    "id": "cccccccc-0000-4000-8000-000000000001",
    "ho_ten": "Nguyễn Văn Học Viên",
    "email": "hocvien.khoakhac@example.com",
}


# ---------------------------------------------------------------- tiện ích in ấn
def _in_tieu_de(txt):
    print("\n" + "=" * 78)
    print(txt)
    print("=" * 78)


def _o(v):
    return "" if v is None else str(v)


def doc_meta(wb):
    """Trả dict metadata từ sheet _HTPLDN_META (đọc thật, không đoán)."""
    if SHEET_META not in wb.sheetnames:
        return {}
    ws = wb[SHEET_META]
    meta = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=2):
        k = _o(row[0].value).strip()
        v = _o(row[1].value).strip() if len(row) > 1 else ""
        if k:
            meta[k] = v
    return meta


def tim_sheet_du_lieu(wb):
    """Sheet dữ liệu = sheet hiển thị đầu tiên khác _HTPLDN_META."""
    for ws in wb.worksheets:
        if ws.title != SHEET_META and ws.sheet_state == "visible":
            return ws
    for ws in wb.worksheets:
        if ws.title != SHEET_META:
            return ws
    return None


def doc_header(ws):
    return [_o(c.value).strip() for c in ws[1]]


def doc_dong_du_lieu(ws, so_cot):
    """Các dòng có cột A (hoc_vien_id) khác rỗng, tính từ dòng 2."""
    ket_qua = []
    for r in range(2, ws.max_row + 1):
        if _o(ws.cell(row=r, column=1).value).strip() == "":
            continue
        ket_qua.append([ws.cell(row=r, column=c).value for c in range(1, so_cot + 1)])
    return ket_qua


def in_cau_truc(duong_dan, nhan="TỆP"):
    """In toàn bộ cấu trúc tệp xlsx — dùng cho --chi-doc và cho bước tự kiểm."""
    p = Path(duong_dan)
    _in_tieu_de(f"{nhan}: {p}")
    print(f"kích thước: {p.stat().st_size} byte")

    # xác nhận là xlsx thật (zip container OOXML), không phải chuỗi chữ đổi đuôi
    try:
        with zipfile.ZipFile(p) as z:
            ten = z.namelist()
            assert "[Content_Types].xml" in ten
        print("định dạng   : ✅ xlsx THẬT (zip OOXML, có [Content_Types].xml)")
    except Exception as e:
        print(f"định dạng   : ❌ KHÔNG phải xlsx hợp lệ — {e}")
        return None

    wb = openpyxl.load_workbook(p)
    print(f"sheet       : {wb.sheetnames}")
    for ws in wb.worksheets:
        print(
            f"\n--- sheet {ws.title!r} · trạng thái={ws.sheet_state} "
            f"· vùng={ws.dimensions} · {ws.max_row} dòng × {ws.max_column} cột"
        )
        if ws.title == SHEET_META:
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=2):
                print(f"    {_o(row[0].value):<18} = {_o(row[1].value)}")
            continue
        print(f"    khóa sheet (protection) : {ws.protection.sheet}")
        print(f"    đóng băng (freeze)      : {ws.freeze_panes}")
        an_cot = [k for k, v in ws.column_dimensions.items() if v.hidden]
        print(f"    cột bị ẩn               : {an_cot or '(không)'}")
        dv = [
            (d.type, d.formula1, str(d.sqref))
            for d in ws.data_validations.dataValidation
        ]
        print(f"    ràng buộc giá trị (dropdown): {dv or '(không)'}")
        print("    nội dung:")
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            print(f"      dòng {row[0].row:>2}: {[_o(c.value) for c in row]}")
    return wb


# ---------------------------------------------------------------- xử lý chính
def main():
    ap = argparse.ArgumentParser(
        description="Điền tệp mẫu điểm danh KTDGKQHT_05 thành fixture 3 hợp lệ + 2 lỗi.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    ap.add_argument("tep_mau", help="Đường dẫn tệp mẫu VỪA TẢI từ nút 'Tải mẫu điểm danh'")
    ap.add_argument("-o", "--ra", help="Đường dẫn tệp fixture ghi ra (.xlsx)")
    ap.add_argument("--chi-doc", action="store_true", help="Chỉ in cấu trúc tệp mẫu, không ghi gì")
    ap.add_argument("--hv-khoa-khac", default=HV_KHOA_KHAC_MAC_DINH["id"],
                    help="hoc_vien_id của học viên thuộc khóa KHÁC (ép ERR-KQ-03)")
    ap.add_argument("--ten-hv-khoa-khac", default=HV_KHOA_KHAC_MAC_DINH["ho_ten"])
    ap.add_argument("--email-hv-khoa-khac", default=HV_KHOA_KHAC_MAC_DINH["email"])
    ap.add_argument("--bo-qua-kiem-cot", action="store_true",
                    help="Vẫn chạy khi bộ cột lệch đặc tả (PHẢI ghi lệch đó vào do/KTDGKQHT_05.md)")
    args = ap.parse_args()

    tep_mau = Path(args.tep_mau).expanduser()
    if not tep_mau.is_file():
        sys.exit(f"❌ Không thấy tệp mẫu: {tep_mau}")

    wb = in_cau_truc(tep_mau, nhan="TỆP MẪU ĐẦU VÀO")
    if wb is None:
        sys.exit("❌ Tệp mẫu không mở được — tải lại từ nút 'Tải mẫu điểm danh'.")

    meta = doc_meta(wb)
    ws = tim_sheet_du_lieu(wb)
    if ws is None:
        sys.exit("❌ Không tìm thấy sheet dữ liệu trong tệp mẫu.")

    _in_tieu_de("KIỂM TỆP MẪU TRƯỚC KHI ĐIỀN")

    # --- metadata (đây là thứ khiến không thể tự dựng tệp) ---
    if not meta:
        print(f"⚠️  KHÔNG có sheet {SHEET_META} → tệp không mang định danh buổi học.")
        print("    Ghi quan sát này vào do/KTDGKQHT_05.md (liên quan vế C3, đặc tả :583).")
    else:
        print(f"metadata tệp ({SHEET_META}):")
        for k, v in meta.items():
            print(f"    {k:<18} = {v}")
        if meta.get("template_type") != "DIEM_DANH":
            sys.exit(
                f"❌ template_type = {meta.get('template_type')!r}, không phải 'DIEM_DANH'.\n"
                "   Đây là tệp mẫu SAI loại (có thể là mẫu điểm kiểm tra). Tải lại ở Tab 4 Điểm danh."
            )
        if not meta.get("lich_hoc_id"):
            print("⚠️  metadata thiếu lich_hoc_id — ghi vào do/ (đặc tả :583 bắt buộc nhúng).")
        else:
            print(f"\n🔴 ĐỐI CHIẾU TAY: lich_hoc_id = {meta['lich_hoc_id']}")
            print("   phải ĐÚNG buổi học bạn vừa chọn trên màn. Lệch → ERR-KQ-09 → FAIL oan.")

    # --- bộ cột ---
    header = doc_header(ws)
    print(f"\nbộ cột đọc được : {header}")
    print(f"bộ cột đặc tả   : {COT_CHUAN}   (srs-fr-03-dao-tao.md:581)")
    khop_cot = header[: len(COT_CHUAN)] == COT_CHUAN
    if khop_cot:
        print("⇒ ✅ khớp đúng 6 cột, đúng thứ tự")
    else:
        print("⇒ ❌ LỆCH đặc tả — đây là một số đo của vế C3, PHẢI ghi vào do/KTDGKQHT_05.md")
        if not args.bo_qua_kiem_cot:
            sys.exit(
                "DỪNG. Ghi lại bộ cột thực tế trước, rồi chạy lại kèm --bo-qua-kiem-cot "
                "nếu vẫn muốn dựng fixture theo bộ cột hiện có."
            )

    # --- dòng dữ liệu ---
    so_cot = max(len(COT_CHUAN), len(header))
    dong = doc_dong_du_lieu(ws, so_cot)
    print(f"\nsố học viên điền sẵn trong tệp mẫu: {len(dong)}")
    for i, d in enumerate(dong, start=1):
        print(f"    HV{i}: {_o(d[0])} · {_o(d[1])} · trạng thái hiện tại={_o(d[4]) or '(trống)'}")

    trong_het = all(_o(d[4]).strip() == "" for d in dong)
    print(f"cột 'Trạng thái điểm danh' rỗng 100%: {'✅ có' if trong_het else '❌ KHÔNG (ghi vào do/)'}")

    if len(dong) < 4:
        sys.exit(
            f"\n❌ Tệp mẫu chỉ có {len(dong)} học viên — cần ≥4 để dựng 3 hợp lệ + 1 dòng 'XYZ'.\n"
            "   Xử lý: vào tab 'Học viên' của khóa, bấm 'Phê duyệt' các đăng ký đang Chờ duyệt\n"
            "   cho đủ ≥4 học viên Đã duyệt, tải lại tệp mẫu, rồi chạy lại script này.\n"
            "   🔴 Duyệt thêm học viên = mutate môi trường chung → PHẢI khai vào do/KTDGKQHT_05.md."
        )
    if len(dong) > 4:
        print(
            f"ℹ️  Tệp mẫu có {len(dong)} học viên; fixture chỉ dùng 4 học viên ĐẦU TIÊN "
            "(các dòng còn lại bị bỏ khỏi fixture để giữ đúng thiết kế đếm 3/2/5)."
        )

    if args.chi_doc:
        _in_tieu_de("CHẾ ĐỘ --chi-doc: không ghi tệp nào.")
        return

    # --- dựng nội dung fixture ---
    def _dong_moi(goc, trang_thai):
        v = list(goc) + [None] * (so_cot - len(goc))
        v[4] = trang_thai
        return v[:so_cot]

    hv_khac = [None] * so_cot
    hv_khac[0] = args.hv_khoa_khac
    hv_khac[1] = args.ten_hv_khoa_khac
    hv_khac[2] = args.email_hv_khoa_khac
    hv_khac[3] = _o(dong[0][3])          # mượn cột Đơn vị cho giống định dạng
    hv_khac[4] = TRANG_THAI_HOP_LE[0]    # trạng thái ĐÚNG enum → lỗi chỉ do học viên khóa khác
    hv_khac[5] = ""

    ra = [
        _dong_moi(dong[0], TRANG_THAI_HOP_LE[0]),   # dòng 2 — hợp lệ
        _dong_moi(dong[1], TRANG_THAI_HOP_LE[1]),   # dòng 3 — hợp lệ
        _dong_moi(dong[2], TRANG_THAI_HOP_LE[2]),   # dòng 4 — hợp lệ
        hv_khac,                                    # dòng 5 — ERR-KQ-03
        _dong_moi(dong[3], TRANG_THAI_SAI),         # dòng 6 — ERR-KQ-04
    ]

    # xóa sạch dòng dữ liệu cũ rồi ghi lại đúng 5 dòng
    ws.delete_rows(2, ws.max_row)
    for i, vals in enumerate(ra, start=2):
        for j, v in enumerate(vals, start=1):
            ws.cell(row=i, column=j, value=v)

    # nới vùng dropdown cho khớp số dòng mới (chỉ để mở bằng Excel cho gọn)
    for d in ws.data_validations.dataValidation:
        try:
            d.sqref = f"E2:E{1 + len(ra)}"
        except Exception:
            pass

    if args.ra:
        tep_ra = Path(args.ra).expanduser()
    else:
        thu_muc = Path(__file__).resolve().parent.parent / "files"
        thu_muc.mkdir(parents=True, exist_ok=True)
        moc = _dt.datetime.now().strftime("%Y%m%d-%H%M")
        tep_ra = thu_muc / f"fixture-KTDGKQHT_05-3hople-2loi-{moc}.xlsx"
    tep_ra.parent.mkdir(parents=True, exist_ok=True)
    wb.save(tep_ra)

    # ================= TỰ KIỂM: mở lại tệp vừa ghi và in toàn bộ =================
    wb2 = in_cau_truc(tep_ra, nhan="TỆP FIXTURE VỪA GHI (mở lại để tự kiểm)")
    if wb2 is None:
        sys.exit("❌ Tệp vừa ghi không mở lại được.")

    meta2 = doc_meta(wb2)
    ws2 = tim_sheet_du_lieu(wb2)
    dong2 = doc_dong_du_lieu(ws2, so_cot)

    _in_tieu_de("BẢNG KỲ VỌNG — đối chiếu với khối xem trước trên màn")
    nhan = ["hợp lệ", "hợp lệ", "hợp lệ", "LỖI ERR-KQ-03 (học viên khóa khác)", "LỖI ERR-KQ-04 (trạng thái XYZ)"]
    print(f"{'dòng':>5} | {'hoc_vien_id':<38} | {'trạng thái':<16} | kỳ vọng")
    print("-" * 110)
    for i, d in enumerate(dong2):
        print(f"{i + 2:>5} | {_o(d[0]):<38} | {_o(d[4]):<16} | {nhan[i] if i < len(nhan) else '?'}")
    print("-" * 110)
    print("Khối xem trước PHẢI ra: Tổng dòng 5 · Hợp lệ 3 · Bỏ qua 0 · Lỗi 2")
    print("Hai dòng lỗi PHẢI có 2 lý do KHÁC NHAU, gắn đúng số dòng 5 và 6.")

    loi = []
    if len(dong2) != 5:
        loi.append(f"số dòng dữ liệu = {len(dong2)}, phải là 5")
    if meta and meta2.get("lich_hoc_id") != meta.get("lich_hoc_id"):
        loi.append("lich_hoc_id trong metadata KHÔNG giữ nguyên sau khi ghi")
    if meta and meta2.get("khoa_hoc_id") != meta.get("khoa_hoc_id"):
        loi.append("khoa_hoc_id trong metadata KHÔNG giữ nguyên sau khi ghi")
    if SHEET_META in wb2.sheetnames and wb2[SHEET_META].sheet_state != wb[SHEET_META].sheet_state:
        loi.append("trạng thái ẩn của sheet metadata bị đổi")
    hop_le_thuc = [_o(d[4]) for d in dong2[:3]]
    if hop_le_thuc != TRANG_THAI_HOP_LE:
        loi.append(f"3 dòng đầu không đúng 3 giá trị hợp lệ: {hop_le_thuc}")
    if _o(dong2[4][4]) != TRANG_THAI_SAI:
        loi.append("dòng 6 không mang giá trị 'XYZ'")
    if _o(dong2[3][0]) != args.hv_khoa_khac:
        loi.append("dòng 5 không mang hoc_vien_id của khóa khác")

    _in_tieu_de("KẾT LUẬN TỰ KIỂM")
    if loi:
        for x in loi:
            print(f"❌ {x}")
        sys.exit("\n❌ Fixture CHƯA dùng được. Sửa xong hãy chạy lại.")
    print("✅ Fixture hợp lệ, là tệp xlsx thật, metadata buổi học giữ nguyên.")
    print(f"✅ Đường dẫn: {tep_ra}")
    print("\n🔴 Trước khi nạp, kiểm tay 2 điều script KHÔNG tự biết:")
    print(f"   1) lich_hoc_id {meta2.get('lich_hoc_id', '(không có)')} đúng buổi đang chọn trên màn.")
    print(f"   2) hoc_vien_id {args.hv_khoa_khac} vẫn TỒN TẠI và vẫn thuộc khóa KHÁC "
          "(nếu đã bị xóa thì vẫn ra ERR-KQ-03 nhưng chỉ chứng minh 'không tồn tại',")
    print("      không chứng minh được ràng buộc 'phải THUỘC khóa' ở đặc tả :592).")


if __name__ == "__main__":
    main()
