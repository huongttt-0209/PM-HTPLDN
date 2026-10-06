from __future__ import annotations

import importlib.util
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "bao-cao-tong-hop-qa"
MASTER = OUT / "Test-case.xlsx"
DEV_VERIFY = ROOT / "output" / "dev-report" / "report-2026-06-07" / "verify-report-2026-06-07.md"

spec = importlib.util.spec_from_file_location("reconcile_full_reports", ROOT / "tools" / "reconcile_full_reports.py")
rec = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(rec)


MODULE_TO_SHEET = {
    "api-consumer": "13. Quản trị hệ thống",
    "audit-log": "13. Quản trị hệ thống",
    "auth": "13. Quản trị hệ thống",
    "bao-cao": "14. Báo cáo thống kê",
    "bieu-mau": "12. Biểu mẫu",
    "cau-hinh-he-thong": "13. Quản trị hệ thống",
    "chi-tra": "09. Chi trả chi phí",
    "chuyen-gia-tvv": "05. Chuyên gia tư vấn",
    "cross-cutting": "13. Quản trị hệ thống",
    "ct-htpldn": "18. Chương trình HTPLDN",
    "danh-gia": "11. Đánh giá",
    "dao-tao-core": "04. Đào tạo tập huấn",
    "dao-tao-hoc-lieu": "04. Đào tạo tập huấn",
    "dashboard": "02. Dashboard tổng quan",
    "doanh-nghiep": "10. Quản lý doanh nghiệp",
    "file-upload": "13. Quản trị hệ thống",
    "hop-dong-tv": "17. Hợp đồng tư vấn",
    "quan-tri/don-vi": "13. Quản trị hệ thống",
    "thong-bao": "13. Quản trị hệ thống",
    "tieu-chi-danh-gia": "11. Đánh giá",
    "tu-van": "16. Tư vấn nhanh",
    "tv-cs": "15. Tư vấn chuyên sâu",
    "tvcs": "15. Tư vấn chuyên sâu",
    "tv-nhanh": "16. Tư vấn nhanh",
    "vu-viec": "08. Vụ việc HTPL",
}


def split_md_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def parse_dev_rows() -> list[dict[str, str]]:
    rows = []
    in_table = False
    for line in DEV_VERIFY.read_text(encoding="utf-8", errors="ignore").splitlines():
        if line.startswith("| #N | slug | module |"):
            in_table = True
            continue
        if in_table and (not line.startswith("|") or line.startswith("|---")):
            if rows:
                break
            continue
        if not in_table or not line.startswith("|"):
            continue
        cells = split_md_row(line)
        if len(cells) < 8 or not cells[0].isdigit():
            continue
        n, slug, module, priority, ba, qa, evidence, note = cells[:8]
        sheet = MODULE_TO_SHEET.get(module)
        if not sheet:
            sheet = "13. Quản trị hệ thống"
        rows.append(
            {
                "n": int(n),
                "id": f"DEVBUG-{int(n):03d}",
                "slug": slug.strip("`"),
                "module": module,
                "sheet": sheet,
                "priority": priority,
                "ba": ba,
                "qa": re.sub(r"<[^>]+>", "", qa),
                "evidence": evidence,
                "note": note,
            }
        )
    return rows


def title_from_slug(slug: str) -> str:
    text = slug.replace("-", " ").replace("_", " ").strip()
    return text[:1].upper() + text[1:]


def scrub_dates(text: str) -> str:
    text = re.sub(r"-(20\d{6})-", "-...-", text)
    text = re.sub(r"\b20\d{6}\b", "...", text)
    text = re.sub(r"\b20\d{2}-\d{2}-\d{2}\b", "...", text)
    return re.sub(r"\b\d{1,2}/\d{1,2}/\d{2,4}\b", "...", text)


def cluster_for(sheet: str, row: dict[str, str]) -> str:
    item = {
        "source_id": row["id"],
            "trace": f"DEV-FIX / {row['module']} / {row['slug']}",
        "title": title_from_slug(row["slug"]),
        "file": "dev-report/report-2026-06-07/verify-report-2026-06-07.md",
    }
    return rec.choose_cluster(sheet, item)


def master_ids(wb) -> set[str]:
    ids = set()
    for ws in wb.worksheets:
        if ws.title == "Tổng hợp":
            continue
        for r in range(4, ws.max_row + 1):
            value = ws.cell(r, 2).value
            if value:
                ids.add(str(value).strip())
    return ids


def dot3_status(row: dict[str, str]) -> tuple[str, str]:
    qa = row["qa"]
    n = row["n"]
    if "PASS" in qa:
        return "PASS", "Đã có report verify dev"
    if "FAIL" in qa:
        return "FAIL", "Report verify dev: bug chưa đạt"
    if n in {14, 69, 74} or "Env" in qa or "BLOCKED" in qa:
        return "CHƯA CHẠY", "Chưa tích hợp"
    if "Chờ BA" in qa:
        return "CHƯA CHẠY", "Chờ BA chốt spec/hướng fix"
    return "CHƯA CHẠY", "Chưa có evidence chạy"


def report_status(report_name: str, row: dict[str, str]) -> tuple[str, str]:
    if report_name == "report-dot-3.xlsx":
        return dot3_status(row)
    if row["n"] in {14, 69, 74}:
        return "CHƯA CHẠY", "Chưa tích hợp"
    return "CHƯA CHẠY", "Chưa có evidence chạy ở đợt này"


def append_master_rows(rows: list[dict[str, str]]) -> dict[str, int]:
    wb = load_workbook(MASTER)
    ids = master_ids(wb)
    added = 0
    by_sheet: dict[str, int] = {}
    for row in rows:
        if row["id"] in ids:
            continue
        ws = wb[row["sheet"]]
        target = ws.max_row + 1
        values = [
            cluster_for(row["sheet"], row),
            row["id"],
            f"DEV-FIX / #{row['n']} / {row['module']} / {row['priority']}",
            title_from_slug(row["slug"]),
            "Bug dev đã được fix và có danh sách verify.",
            f"BA={row['ba']}; QA={row['qa']}",
            f"1. Re-verify bug #{row['n']} theo dev-report.\n2. Đối chiếu evidence/note trong verify report.",
            scrub_dates(f"{row['evidence']}. {row['note']}").strip(),
        ]
        ws.append(values)
        rec.copy_row_style(ws, max(4, target - 1), target)
        ids.add(row["id"])
        added += 1
        by_sheet[row["sheet"]] = by_sheet.get(row["sheet"], 0) + 1
    rec.update_master_summary(wb)
    wb.save(MASTER)
    return {"added": added, **by_sheet}


def sync_reports(rows: list[dict[str, str]]) -> dict[str, dict[str, int]]:
    master = load_workbook(MASTER, data_only=False)
    out: dict[str, dict[str, int]] = {}
    for report_name, label in rec.REPORTS:
        wb = load_workbook(OUT / report_name)
        added = 0
        updated = 0
        for row in rows:
            sheet = row["sheet"]
            ws = wb[sheet]
            mws = master[sheet]
            rec.ensure_report_columns(ws)
            existing = {str(ws.cell(r, 2).value).strip(): r for r in range(4, ws.max_row + 1) if ws.cell(r, 2).value}
            if row["id"] not in existing:
                source_row = None
                for r in range(4, mws.max_row + 1):
                    if str(mws.cell(r, 2).value or "").strip() == row["id"]:
                        source_row = r
                        break
                if source_row is None:
                    continue
                target = ws.max_row + 1
                for col in range(1, 9):
                    rec.copy_cell(mws.cell(source_row, col), ws.cell(target, col))
                    ws.cell(target, col).alignment = Alignment(wrap_text=True, vertical="top")
                existing[row["id"]] = target
                added += 1
            target = existing[row["id"]]
            status, reason = report_status(report_name, row)
            if ws.cell(target, 9).value != status or ws.cell(target, 10).value != reason:
                ws.cell(target, 9).value = status
                ws.cell(target, 10).value = reason
                rec.set_status_style(ws.cell(target, 9), status)
                ws.cell(target, 10).alignment = Alignment(wrap_text=True, vertical="top")
                updated += 1
        totals = rec.update_report_summary(wb, label)
        wb.save(OUT / report_name)
        out[report_name] = {"added": added, "updated": updated, **totals}
    return out


def main() -> None:
    rows = parse_dev_rows()
    print({"source_dev_rows": len(rows), "master": append_master_rows(rows), "reports": sync_reports(rows)})


if __name__ == "__main__":
    main()
