from __future__ import annotations

import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "bao-cao-tong-hop-qa"
XLSX = OUT_DIR / "Test-case.xlsx"
MD_OUT = OUT_DIR / "BAO-CAO-UAT-VERIFY-BUG-3-DOT-GAN-NHAT-2026-06.md"
UAT1 = ROOT / "output" / "BA-report" / "Bug-report" / "2026-06-02" / "verify-report-2026-06-02.md"


SHEET_BY_MODULE = {
    "bao-cao": "14. Báo cáo thống kê",
    "bm": "12. Biểu mẫu",
    "chi-tra": "09. Chi trả chi phí",
    "ct-htpldn": "18. Chương trình HTPLDN",
    "danh-gia": "11. Đánh giá",
    "dao-tao": "04. Đào tạo tập huấn",
    "doanh-nghiep": "10. Quản lý doanh nghiệp",
    "hoi-dap": "03. Hỏi đáp pháp lý",
    "qtht-cau-hinh-ht": "13. Quản trị hệ thống",
    "qtht-danh-muc": "13. Quản trị hệ thống",
    "qtht-nhat-ky": "13. Quản trị hệ thống",
    "qtht-tai-khoan": "13. Quản trị hệ thống",
    "qtht-vai-tro": "13. Quản trị hệ thống",
    "tu-van-chuyen-sau": "15. Tư vấn chuyên sâu",
    "tu-van-nhanh": "16. Tư vấn nhanh",
    "tu-van-vien-cg": "05. Chuyên gia tư vấn",
    "vu-viec": "08. Vụ việc HTPL",
}


UAT2_ROWS = [
    ("22", "dao-tao", "Mô tả bài giảng tùy chọn theo chốt BA"),
    ("56", "dao-tao", "Khóa học có mô tả công khai dạng rich editor"),
    ("57", "tu-van-chuyen-sau", "Tắt spellcheck toàn cục cho các ô nhập/tìm kiếm"),
    ("58", "tu-van-nhanh", "Điểm phù hợp tư vấn nhanh chuẩn hóa 0-100%"),
    ("59", "hoi-dap", "Chọn và bỏ chọn câu hỏi gợi ý khi phân công xử lý"),
    ("60", "tu-van-chuyen-sau", "Form sửa TVCS có đủ nhóm thông tin cơ bản"),
    ("63", "tu-van-chuyen-sau", "Chuyên gia xem tư liệu pháp lý và ghi audit"),
    ("64", "tu-van-chuyen-sau", "Danh sách TVCS chỉ có xem chi tiết, duyệt ở chi tiết"),
    ("66", "dao-tao", "Bài giảng có cột/bộ lọc công khai và thông tin ảnh"),
    ("68", "tu-van-chuyen-sau", "Đổi Tóm tắt thành Tiêu đề và bắt buộc nhập"),
    ("69", "qtht-danh-muc", "Danh mục chỉ cột Tên có sắp xếp"),
    ("72", "qtht-danh-muc", "Danh mục VSIC vẫn hiển thị"),
    ("80", "qtht-tai-khoan", "Email tạo tài khoản gửi link kích hoạt, không gửi mật khẩu thô"),
    ("85", "bao-cao", "Excel xuất báo cáo có header và kẻ viền"),
    ("96", "bao-cao", "Báo cáo dùng Đơn vị, bỏ Địa bàn/NHT theo chốt BA"),
]


def clean_cell(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"<[^>]+>", "", text)
    text = text.replace("**", "").replace("`", "").replace("~~", "")
    return " ".join(text.split())


def parse_uat1_rows() -> list[tuple[str, str, str, str]]:
    rows: list[tuple[str, str, str, str]] = []
    for line in UAT1.read_text(encoding="utf-8").splitlines():
        if not line.startswith("|"):
            continue
        cells = [clean_cell(c.strip()) for c in line.strip().strip("|").split("|")]
        if len(cells) < 7 or not cells[0].isdigit():
            continue
        stt, module, title, qa = cells[0], cells[1], cells[2], cells[4]
        rows.append((stt, module, title, qa))
    return rows


def existing_ids(wb) -> set[str]:
    ids: set[str] = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=4, max_col=2, values_only=True):
            if row and row[1]:
                ids.add(str(row[1]).strip())
    return ids


def append_tc(ws, row_values: list[str]) -> None:
    row_idx = ws.max_row + 1
    ws.append(row_values)
    template_row = max(4, row_idx - 1)
    for col in range(1, 9):
        src = ws.cell(template_row, col)
        dst = ws.cell(row_idx, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = Alignment(wrap_text=True, vertical="top")


def recreate_sheet(wb, title: str, headers: list[str], rows: list[list[str]]) -> None:
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [16, 18, 22, 14, 14, 14, 18, 60]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width
    ws.freeze_panes = "A2"


def update_summary(wb) -> None:
    ws = wb["Tổng hợp"]
    total = 0
    for row in range(2, ws.max_row + 1):
        module = ws.cell(row, 2).value
        if not module:
            continue
        if str(module).strip().lower() == "17 module":
            continue
        sheet = next((name for name in wb.sheetnames if name.endswith(str(module))), None)
        if not sheet:
            continue
        count = sum(
            1
            for values in wb[sheet].iter_rows(min_row=4, max_col=2, values_only=True)
            if values[1]
        )
        total += count
        ws.cell(row, 3).value = count
        ws.cell(row, 4).value = count
        ws.cell(row, 5).value = "Đã bổ sung TC UAT/verify bug; đợt cuối PASS 100% trong phạm vi đối tác."
    for row in range(2, ws.max_row + 1):
        module = ws.cell(row, 2).value
        if str(module).strip().lower() == "17 module":
            ws.cell(row, 3).value = total
            ws.cell(row, 4).value = total
            ws.cell(row, 5).value = "Tổng sau khi bổ sung TC UAT/verify bug."


def write_markdown(total_added: int, added_by_sheet: dict[str, int]) -> None:
    by_sheet_rows = "\n".join(
        f"| {sheet} | {count} |" for sheet, count in sorted(added_by_sheet.items()) if count
    )
    content = f"""# BÁO CÁO UAT / VERIFY BUG — 3 ĐỢT GẦN NHẤT

| Thông tin | Giá trị |
|---|---|
| Phạm vi | UAT + verify bug đối tác tháng 06/2026 |
| Nguồn | `verify-report-2026-06-02.md`; `bug-report-verify-uat-dot2-2026-06-04.md` |
| File testcase cập nhật | `output/bao-cao-tong-hop-qa/Test-case.xlsx` |
| Ngày lập | 2026-06-25 |

## 1. Tổng hợp 3 đợt

| Đợt | Ngày snapshot | Phạm vi | Tổng mục | PASS | FAIL/PARTIAL | Tỷ lệ PASS | Kết luận |
|---|---|---|---:|---:|---:|---:|---|
| Đợt 1 | 2026-06-02 | Verify 92 bug theo snapshot 01/06 | 92 | 88 | 4 | 95,65% | Còn lỗi cần fix/retest |
| Đợt 2 | 2026-06-04 / retest 2026-06-05 | Verify UAT đợt 2 theo BA confirm | 15 | 8 | 7 | 53,33% | Còn 7 mục open sau R-verify-3 |
| Đợt 3 | 2026-06-07 18:58:58 | Retest R-verify-4 sau dev deploy | 15 | 15 | 0 | 100% | PASS toàn bộ phạm vi đối tác |

## 2. Testcase đã bổ sung

Đã bổ sung {total_added} testcase mới vào workbook, dùng ID `UAT1-STT-xxx` và `UAT2-STT-xxx`. Script chỉ thêm khi ID chưa tồn tại để tránh trùng.

| Sheet module | Số TC bổ sung |
|---|---:|
{by_sheet_rows}

## 3. Ghi chú phạm vi

Các mục hạ tầng/chờ BA hoặc verify nội bộ không thuộc bảng PASS/FAIL đối tác được tách khỏi phạm vi báo cáo chính. Theo yêu cầu mới nhất, không thực hiện reverify live bổ sung trong ngày 2026-06-25.
"""
    MD_OUT.write_text(content, encoding="utf-8")


def main() -> None:
    wb = load_workbook(XLSX)
    ids = existing_ids(wb)
    added_by_sheet: dict[str, int] = {}

    for stt, module, title, qa in parse_uat1_rows():
        sheet = SHEET_BY_MODULE.get(module)
        if not sheet or sheet not in wb.sheetnames:
            continue
        tc_id = f"UAT1-STT-{int(stt):03d}"
        if tc_id in ids:
            continue
        append_tc(
            wb[sheet],
            [
                "UAT/Verify bug đối tác 2026-06-02",
                tc_id,
                f"UAT1-STT-{stt}",
                title,
                "Có tài khoản/role phù hợp và dữ liệu test theo bug đối tác.",
                f"STT {stt}",
                "Mở chức năng liên quan, thực hiện thao tác theo mô tả bug và kiểm tra hành vi sau fix.",
                "Hệ thống xử lý đúng theo SRS/BA confirm; không còn lỗi đã ghi nhận.",
            ],
        )
        ids.add(tc_id)
        added_by_sheet[sheet] = added_by_sheet.get(sheet, 0) + 1

    for stt, module, title in UAT2_ROWS:
        sheet = SHEET_BY_MODULE.get(module)
        if not sheet or sheet not in wb.sheetnames:
            continue
        tc_id = f"UAT2-STT-{int(stt):03d}"
        if tc_id in ids:
            continue
        append_tc(
            wb[sheet],
            [
                "UAT/Verify bug đối tác 2026-06-04 đến 2026-06-07",
                tc_id,
                f"UAT2-STT-{stt}",
                title,
                "Có tài khoản/role phù hợp và dữ liệu theo quyết định BA.",
                f"STT {stt}",
                "Thực hiện thao tác theo BA confirm và kiểm tra lại trên UI/API/export nếu áp dụng.",
                "Kết quả đúng theo BA confirm; snapshot R-verify-4 ngày 07/06/2026 PASS.",
            ],
        )
        ids.add(tc_id)
        added_by_sheet[sheet] = added_by_sheet.get(sheet, 0) + 1

    report_rows = [
        ["Đợt 1", "2026-06-02", "Verify 92 bug theo snapshot 01/06", 92, 88, 4, "95,65%", "Còn lỗi cần fix/retest"],
        ["Đợt 2", "2026-06-04 / 2026-06-05", "Verify UAT đợt 2 theo BA confirm", 15, 8, 7, "53,33%", "Còn 7 mục open sau R-verify-3"],
        ["Đợt 3", "2026-06-07 18:58:58", "Retest R-verify-4 sau dev deploy", 15, 15, 0, "100%", "PASS toàn bộ phạm vi đối tác"],
    ]
    recreate_sheet(
        wb,
        "Report 3 dot UAT",
        ["Đợt", "Ngày", "Phạm vi", "Tổng mục", "PASS", "FAIL/PARTIAL", "Tỷ lệ PASS", "Kết luận"],
        report_rows,
    )
    recreate_sheet(
        wb,
        "Ngoai pham vi",
        ["Mã", "Nhóm", "Lý do", "Xử lý trong report", "Ghi chú"],
        [
            ["INFRA", "Hạ tầng", "Cổng PLQG/mTLS hoặc sandbox chưa sẵn sàng", "Không tính vào mẫu đối tác", "Theo dõi riêng khi có môi trường"],
            ["BA", "Chờ BA/spec", "Cần BA chốt hướng xử lý hoặc cập nhật SRS", "Không tính vào mẫu đối tác", "Không đưa vào tập TC gửi đối tác"],
            ["INTERNAL-REVERIFY", "Verify nội bộ", "Không chạy reverify live ngày 25/06 theo yêu cầu mới nhất", "Không tính vào bảng 3 đợt", "Dùng báo cáo nguồn đã có"],
        ],
    )
    update_summary(wb)
    wb.save(XLSX)
    write_markdown(sum(added_by_sheet.values()), added_by_sheet)
    print({"added": sum(added_by_sheet.values()), "by_sheet": added_by_sheet, "xlsx": str(XLSX), "markdown": str(MD_OUT)})


if __name__ == "__main__":
    main()
