from __future__ import annotations

import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment


ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "output" / "bao-cao-tong-hop-qa" / "Test-case.xlsx"
SOURCE = ROOT / "output" / "BA-report" / "Bug-report" / "2026-06-01" / "todo-uat-2026-05-26-v3-aligned-xlsx.md"

MODULE_MAP = {
    "0. Chung": "02. Dashboard tổng quan",
    "II Hỏi đáp": "03. Hỏi đáp pháp lý",
    "III Đào tạo": "04. Đào tạo tập huấn",
    "III Video": "04. Đào tạo tập huấn",
    "IV TVV": "05. Chuyên gia tư vấn",
    "IV TC TV": "07. Tổ chức tư vấn",
    "V VV": "08. Vụ việc HTPL",
    "V DN": "10. Quản lý doanh nghiệp",
    "VI Đánh giá": "11. Đánh giá",
    "VII Biểu mẫu": "12. Biểu mẫu",
    "X TVCS": "15. Tư vấn chuyên sâu",
    "XI KHTH": "18. Chương trình HTPLDN",
    "XI Đợt BC": "18. Chương trình HTPLDN",
}

SUMMARY_TO_SHEET = {
    "Dashboard tổng quan": "02. Dashboard tổng quan",
    "Hỏi đáp pháp lý": "03. Hỏi đáp pháp lý",
    "Đào tạo tập huấn": "04. Đào tạo tập huấn",
    "Chuyên gia tư vấn": "05. Chuyên gia tư vấn",
    "Người hỗ trợ pháp luật": "06. Người hỗ trợ pháp luật",
    "Tổ chức tư vấn": "07. Tổ chức tư vấn",
    "Vụ việc HTPL": "08. Vụ việc HTPL",
    "Chi trả chi phí": "09. Chi trả chi phí",
    "Quản lý doanh nghiệp": "10. Quản lý doanh nghiệp",
    "Đánh giá": "11. Đánh giá",
    "Biểu mẫu": "12. Biểu mẫu",
    "Quản trị hệ thống": "13. Quản trị hệ thống",
    "Báo cáo thống kê": "14. Báo cáo thống kê",
    "Tư vấn chuyên sâu": "15. Tư vấn chuyên sâu",
    "Tư vấn nhanh": "16. Tư vấn nhanh",
    "Hợp đồng tư vấn": "17. Hợp đồng tư vấn",
    "Chương trình HTPLDN": "18. Chương trình HTPLDN",
}


def clean(text: str) -> str:
    return " ".join(re.sub(r"\*\*|`", "", text).strip().split())


def parse_partner_rows() -> list[dict[str, str | int]]:
    rows = []
    in_table = False
    for line in SOURCE.read_text(encoding="utf-8").splitlines():
        if line.startswith("| STT | Row |"):
            in_table = True
            continue
        if in_table and line.startswith("|---"):
            continue
        if in_table and not line.startswith("|"):
            break
        if not in_table or not line.startswith("|"):
            continue
        cells = [clean(c) for c in line.strip().strip("|").split("|")]
        if len(cells) >= 10 and cells[0].isdigit():
            rows.append(
                {
                    "stt": int(cells[0]),
                    "row": cells[1],
                    "type": cells[2],
                    "priority": cells[3],
                    "module": cells[5],
                    "title": cells[6],
                    "status": cells[7],
                    "bug": cells[8],
                    "note": cells[9],
                }
            )
    if len(rows) != 55:
        raise RuntimeError(f"Expected 55 partner rows, parsed {len(rows)}")
    return rows


def copy_row_style(ws, source_row: int, target_row: int) -> None:
    for col in range(1, 9):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = Alignment(wrap_text=True, vertical="top")
        dst.number_format = src.number_format


def existing_ids(wb) -> set[str]:
    ids = set()
    for ws in wb.worksheets:
        if ws.title == "Tổng hợp":
            continue
        for row in ws.iter_rows(min_row=4, max_col=2, values_only=True):
            if row[1]:
                ids.add(str(row[1]).strip())
    return ids


def update_summary(wb) -> None:
    ws = wb["Tổng hợp"]
    total = 0
    for row in range(2, ws.max_row + 1):
        module = ws.cell(row, 2).value
        if not module or str(module).strip().lower() == "17 module":
            continue
        sheet = SUMMARY_TO_SHEET.get(str(module))
        if not sheet:
            continue
        count = sum(
            1
            for values in wb[sheet].iter_rows(min_row=4, max_col=2, values_only=True)
            if values[1]
        )
        total += count
        ws.cell(row, 3).value = count
        ws.cell(row, 4).value = count
        ws.cell(row, 5).value = "Đã bổ sung TC UAT/verify bug và task đối tác; đợt cuối PASS trong phạm vi sau loại hạ tầng."
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 2).value).strip().lower() == "17 module":
            ws.cell(row, 3).value = total
            ws.cell(row, 4).value = total
            ws.cell(row, 5).value = "Tổng sau khi bổ sung TC UAT/verify bug/task đối tác."


def main() -> None:
    wb = load_workbook(XLSX)
    ids = existing_ids(wb)
    added = []
    for item in parse_partner_rows():
        stt = int(item["stt"])
        if f"UAT1-STT-{stt:03d}" in ids or f"UAT2-STT-{stt:03d}" in ids or f"DTAC-STT-{stt:03d}" in ids:
            continue
        sheet = MODULE_MAP[str(item["module"])]
        ws = wb[sheet]
        row = ws.max_row + 1
        ws.append(
            [
                "Task/Bug đối tác",
                f"DTAC-STT-{stt:03d}",
                f"{item['row']} | {item['type']} | {item['priority']}",
                item["title"],
                "Có tài khoản/role và dữ liệu test theo danh sách đối tác.",
                f"STT {stt} - {item['module']}",
                "Thực hiện theo mô tả/mong muốn của đối tác; kiểm tra UI/API/export nếu áp dụng.",
                item["note"],
            ]
        )
        copy_row_style(ws, max(4, row - 1), row)
        ids.add(f"DTAC-STT-{stt:03d}")
        added.append((sheet, f"DTAC-STT-{stt:03d}", item["title"]))
    update_summary(wb)
    wb.save(XLSX)
    print({"added": len(added), "items": added})


if __name__ == "__main__":
    main()
