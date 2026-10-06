#!/usr/bin/env python3
import json
import time
import urllib.error
import urllib.request
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


BASE = "http://103.172.236.130:3000/api/v1"
WORKBOOK = Path("output/bao-cao-tong-hop-qa/report-dot-3.xlsx")
OUT = Path("output/bao-cao-tong-hop-qa/m05-c1-permission-2026-06-26")
EVIDENCE = OUT / "evidence.jsonl"
SUMMARY = OUT / "summary.md"
PASSWORD = "Secret@123"
OTP = "666666"
KNOWN_TW_ACTIVE_ID = "978354d7-feac-4330-a750-6b8c07b46c24"


def http(method, path, token=None, data=None):
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    body = json.dumps(data, ensure_ascii=False).encode() if data is not None else None
    req = urllib.request.Request(BASE + path, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            raw = resp.read().decode()
            return resp.status, json.loads(raw) if raw else None
    except urllib.error.HTTPError as exc:
        raw = exc.read().decode()
        try:
            payload = json.loads(raw)
        except Exception:
            payload = raw
        return exc.code, payload


def login(username, pause_before=75):
    time.sleep(pause_before)
    status, payload = http("POST", "/auth/login", data={"username": username, "password": PASSWORD})
    if status != 200 or not isinstance(payload, dict) or not payload.get("success"):
        return None, {"step": "login", "status": status, "payload": payload}
    time.sleep(2)
    status, payload = http("POST", "/auth/verify-otp", data={"otpToken": payload["data"]["otpToken"], "otpCode": OTP})
    if status != 200 or not isinstance(payload, dict) or not payload.get("success"):
        return None, {"step": "otp", "status": status, "payload": payload}
    token = payload["data"]["accessToken"]
    return token, None


def slim(payload):
    if isinstance(payload, dict) and "error" in payload:
        return {"success": payload.get("success"), "error": payload.get("error")}
    if isinstance(payload, dict):
        data = payload.get("data")
        if isinstance(data, dict):
            keep = {k: data.get(k) for k in (
                "id", "maTvv", "hoTen", "donViId", "trangThai", "version", "email", "taiKhoanId"
            ) if k in data}
            return {"success": payload.get("success"), "data": keep, "meta": payload.get("meta")}
        if isinstance(data, list):
            sample = []
            for item in data[:3]:
                sample.append({k: item.get(k) for k in ("id", "maTvv", "hoTen", "donViId", "trangThai") if k in item})
            return {"success": payload.get("success"), "count": len(data), "sample": sample, "meta": payload.get("meta")}
    return payload


def call(method, path, token=None, data=None):
    status, payload = http(method, path, token=token, data=data)
    return {"method": method, "path": path, "request": data, "status": status, "response": slim(payload), "raw": payload}


def record(results, case_id, verdict, note, calls):
    rec = {"caseId": case_id, "verdict": verdict, "note": note, "calls": calls}
    results.append(rec)
    with EVIDENCE.open("a", encoding="utf-8") as f:
        f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    print(case_id, verdict, note)


def update_workbook(results):
    wb = load_workbook(WORKBOOK)
    ws = wb["05. Chuyên gia tư vấn"]
    header_row = None
    for row in range(1, 20):
        vals = [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
        if "Kết quả" in vals:
            header_row = row
            break
    headers = [ws.cell(header_row, col).value for col in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}
    by_id = {r["caseId"]: r for r in results if r["verdict"] in ("PASS", "FAIL")}
    updated = []
    for row in range(header_row + 1, ws.max_row + 1):
        case_id = ws.cell(row, idx["ID"]).value
        if case_id in by_id and ws.cell(row, idx["Kết quả"]).value == "CHƯA CHẠY" and not ws.cell(row, idx["Lý do"]).value:
            ws.cell(row, idx["Kết quả"]).value = by_id[case_id]["verdict"]
            ws.cell(row, idx["Lý do"]).value = None
            updated.append((case_id, by_id[case_id]["verdict"]))

    counts = Counter()
    reasons = Counter()
    for row in range(header_row + 1, ws.max_row + 1):
        if not any(ws.cell(row, c).value is not None for c in range(1, ws.max_column + 1)):
            continue
        res = ws.cell(row, idx["Kết quả"]).value
        reason = ws.cell(row, idx["Lý do"]).value
        counts[res] += 1
        if res == "CHƯA CHẠY":
            reasons[reason] += 1

    summary_ws = wb["Tổng hợp"]
    for row in range(1, summary_ws.max_row + 1):
        module = summary_ws.cell(row, 2).value
        if module and "Chuyên gia tư vấn" in str(module):
            total = sum(counts.values())
            summary_ws.cell(row, 3).value = total
            summary_ws.cell(row, 4).value = counts["PASS"]
            summary_ws.cell(row, 5).value = (
                f"Report đợt 3; cập nhật QA bổ sung 2026-06-26 C1 permission; "
                f"PASS {counts['PASS']}/{total}; FAIL {counts['FAIL']}; "
                f"CHƯA CHẠY {counts['CHƯA CHẠY']} "
                f"(Lý do trống {reasons[None]}; Chưa tích hợp {reasons['Chưa tích hợp']})."
            )
            break
    wb.save(WORKBOOK)
    return updated, counts, reasons


def first_same_unit_active(token):
    known = call("GET", f"/tu-van-viens/{KNOWN_TW_ACTIVE_ID}", token)
    data = (known["raw"] or {}).get("data") or {}
    if known["status"] == 200 and data.get("trangThai") == "HOAT_DONG":
        return data, [known]
    base = call("GET", "/tu-van-viens?page=1&pageSize=100&trangThai=HOAT_DONG", token)
    for item in (base["raw"] or {}).get("data") or []:
        detail = call("GET", f"/tu-van-viens/{item['id']}", token)
        data = (detail["raw"] or {}).get("data") or {}
        if detail["status"] == 200 and data.get("donViId") == "00000000-0000-4000-8000-000000000001":
            return data, [base, detail]
    return None, [base]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    EVIDENCE.write_text("", encoding="utf-8")
    results = []

    tokens = {}
    login_notes = []
    for username in ["cb_nv_tw_05", "cb_nv_dp_03", "cb_nv_bn_03", "cb_pd_tw_03", "nht_01", "9999999990"]:
        token, err = login(username)
        if token:
            tokens[username] = token
            login_notes.append(f"{username}: OK")
        else:
            login_notes.append(f"{username}: BLOCKED {err}")

    if "cb_nv_tw_05" not in tokens:
        SUMMARY.write_text("# C1 Permission\n\nBlocked: cannot login `cb_nv_tw_05`.\n", encoding="utf-8")
        print("BLOCKED no cb_nv_tw_05")
        return

    tw_active, baseline_calls = first_same_unit_active(tokens["cb_nv_tw_05"])
    moi_call = call("GET", "/tu-van-viens?page=1&pageSize=20&trangThai=MOI_DANG_KY", tokens["cb_nv_tw_05"])
    tw_moi = ((moi_call["raw"] or {}).get("data") or [None])[0]

    if tw_active and "cb_nv_dp_03" in tokens:
        patch = call("PATCH", f"/tu-van-viens/{tw_active['id']}", tokens["cb_nv_dp_03"], {
            "diaChi": "QA C1 DP tamper", "version": tw_active["version"]
        })
        verdict = "PASS" if patch["status"] in (403, 404) else "FAIL"
        calls = baseline_calls + [patch]
        record(results, "TC-CG-022", verdict, "CB_NV_DP sửa TVV TW/khác đơn vị bị chặn", calls)
        record(results, "TC-CG-307", verdict, "CB_NV_DP cập nhật năng lực TVV ngoài đơn vị bị chặn", calls)
        record(results, "TC-CG-PERM-004", verdict, "IDOR: CB_NV_DP sửa TVV ngoài đơn vị bị chặn", calls)

        if tw_moi:
            td = call("POST", f"/tu-van-viens/{tw_moi['id']}/tham-dinh", tokens["cb_nv_dp_03"], {
                "nhom1KetQua": True, "nhom2Diem": 4, "nhom3Diem": None, "nhom4ThamGia": False,
                "ketLuan": "DAT", "trinhDuyet": True, "version": tw_moi["version"]
            })
            record(results, "TC-CG-508", "PASS" if td["status"] in (403, 404) else "FAIL", "CB_NV_DP thẩm định TVV TW bị chặn", [moi_call, td])

        dp_list = call("GET", "/tu-van-viens?page=1&pageSize=100", tokens["cb_nv_dp_03"])
        contains_btp = any(str(x.get("maTvv", "")).startswith("TVV-BTP-TW") for x in ((dp_list["raw"] or {}).get("data") or []))
        record(results, "TC-CG-PERM-003", "PASS" if dp_list["status"] == 200 and not contains_btp else "FAIL", "CB_NV_DP list không chứa TVV BTP-TW ngoài đơn vị", [dp_list])

    if tw_active and "cb_nv_bn_03" in tokens:
        bn_list = call("GET", "/tu-van-viens?page=1&pageSize=100", tokens["cb_nv_bn_03"])
        detail = call("GET", f"/tu-van-viens/{tw_active['id']}", tokens["cb_nv_bn_03"])
        contains_btp = any(str(x.get("maTvv", "")).startswith("TVV-BTP-TW") for x in ((bn_list["raw"] or {}).get("data") or []))
        verdict = "PASS" if bn_list["status"] == 200 and not contains_btp and detail["status"] in (403, 404) else "FAIL"
        record(results, "TC-CG-PERM-002", verdict, "CB_NV_BN không xem data BTP-TW khác scope", [bn_list, detail])

    if tw_active and "nht_01" in tokens:
        patch = call("PATCH", f"/tu-van-viens/{tw_active['id']}", tokens["nht_01"], {
            "diaChi": "QA C1 NHT tamper", "version": tw_active["version"]
        })
        verdict = "PASS" if patch["status"] in (403, 404) else "FAIL"
        record(results, "TC-CG-304", verdict, "NHT cập nhật hồ sơ TVV khác đơn vị bị chặn", [patch])
        record(results, "TC-CG-1005", verdict, "NHT cố cập nhật hồ sơ khác bị chặn", [patch])

    if tw_active and "9999999990" in tokens:
        list_call = call("GET", "/tu-van-viens?page=1&pageSize=10", tokens["9999999990"])
        detail = call("GET", f"/tu-van-viens/{tw_active['id']}", tokens["9999999990"])
        verdict = "PASS" if list_call["status"] in (403, 404) and detail["status"] in (403, 404) else "FAIL"
        record(results, "TC-CG-707", verdict, "DN/user không có quyền MLTV backend bị chặn", [list_call, detail])

    updated, counts, reasons = update_workbook(results)
    SUMMARY.write_text(
        "# C1 Permission\n\n"
        + "\n".join(f"- {x}" for x in login_notes)
        + "\n\n"
        + f"- Updated: {updated}\n"
        + f"- Counts: {dict(counts)}\n"
        + f"- Reasons: {dict(reasons)}\n"
        + f"- Evidence: `{EVIDENCE}`\n",
        encoding="utf-8",
    )
    print(json.dumps({"updated": updated, "counts": dict(counts), "reasons": {str(k): v for k, v in reasons.items()}}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
