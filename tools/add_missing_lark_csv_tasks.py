from __future__ import annotations

import csv
import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment


ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "output" / "bao-cao-tong-hop-qa" / "Test-case.xlsx"
SOURCE = (
    ROOT
    / "output"
    / "BA-report"
    / "Bug-report"
    / "2026-06-04"
    / "PM_HTPLDN_Danh_sách_tối_ưu_Bug+API_Phần_mềm.csv"
)

MODULE_MAP = {
    "0. Chung": "02. Dashboard tổng quan",
    "II. Quản lý tiếp nhận hỏi đáp vướng mắc pháp lý": "03. Hỏi đáp pháp lý",
    "III. Quản lý thông tin chương trình đào tạo tập huấn": "04. Đào tạo tập huấn",
    "IV. Quản lý thông tin Tư vấn viên, cộng tác viên tư vấn": "05. Chuyên gia tư vấn",
    "V. Quản lý thông tin Hỗ trợ pháp lý theo vụ việc": "08. Vụ việc HTPL",
    "VI. Quản lý thông tin Đánh giá chất lượng vụ việc hỗ trợ pháp lý": "11. Đánh giá",
    "VII. Quản lý biểu mẫu": "12. Biểu mẫu",
    "VIII. Quản trị hệ thống": "13. Quản trị hệ thống",
    "IX. Báo cáo thống kê": "14. Báo cáo thống kê",
    "X. Quản lý tư vấn chuyên sâu với chuyên gia": "15. Tư vấn chuyên sâu",
    "XI. Quản lý thông tin kế hoạch tổng hợp về chương trình hỗ trợ pháp lý cho doanh nghiệp": "18. Chương trình HTPLDN",
}

SUMMARY_TO_SHEET = {
    "Dashboard tổng quan": "02. Dashboard tổng quan",
    "Hỏi đáp pháp lý": "03. Hỏi đáp pháp lý",
    "Đào tạo tập huấn": "04. Đào tạo tập huấn",
    "Chuyên gia tư vấn": "05. Chuyên gia tư vấn",
    "Người hỗ trợ pháp luật": "06. Người hỗ trợ pháp luật",
    "Tổ chức tư vấn": "07. Tổ chức tư vấn",
    "Vụ việc HTPL": "08. Vụ việc HTPL",
    "Chi trả chi phí": "09. Chi trả chi phí",
    "Quản lý doanh nghiệp": "10. Quản lý doanh nghiệp",
    "Đánh giá": "11. Đánh giá",
    "Biểu mẫu": "12. Biểu mẫu",
    "Quản trị hệ thống": "13. Quản trị hệ thống",
    "Báo cáo thống kê": "14. Báo cáo thống kê",
    "Tư vấn chuyên sâu": "15. Tư vấn chuyên sâu",
    "Tư vấn nhanh": "16. Tư vấn nhanh",
    "Hợp đồng tư vấn": "17. Hợp đồng tư vấn",
    "Chương trình HTPLDN": "18. Chương trình HTPLDN",
}


def parse_rows() -> list[dict[str, str | int]]:
    raw = list(csv.reader(SOURCE.open(encoding="utf-8-sig")))
    header_row = next(i for i, row in enumerate(raw) if row and row[0].strip() == "STT")
    headers = [h.strip() for h in raw[header_row]]
    rows = []
    for row in raw[header_row + 1 :]:
        if not row or not row[0].strip().isdigit():
            continue
        item = {
            headers[i] or f"col{i}": (row[i].strip() if i < len(row) else "")
            for i in range(len(headers))
        }
        if item.get("Phân loại") not in {"Bug", "Task"}:
            continue
        item["stt"] = int(item["STT"])
        rows.append(item)
    return rows


def copy_row_style(ws, source_row: int, target_row: int) -> None:
    for col in range(1, 9):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = Alignment(wrap_text=True, vertical="top")
        dst.number_format = src.number_format


def existing_partner_stt(wb) -> set[int]:
    found = set()
    for ws in wb.worksheets:
        if ws.title == "Tổng hợp":
            continue
        for values in ws.iter_rows(min_row=4, min_col=2, max_col=2, values_only=True):
            value = values[0]
            if not value:
                continue
            match = re.match(r"^(?:UAT1|UAT2|DTAC)-STT-(\d+)", str(value).strip())
            if match:
                found.add(int(match.group(1)))
    return found


def update_summary(wb) -> None:
    ws = wb["Tổng hợp"]
    total = 0
    for row in range(2, ws.max_row + 1):
        module = ws.cell(row, 2).value
        if not module or str(module).strip().lower() == "17 module":
            continue
        sheet = SUMMARY_TO_SHEET.get(str(module))
        if not sheet:
            continue
        count = sum(1 for values in wb[sheet].iter_rows(min_row=4, min_col=2, max_col=2, values_only=True) if values[0])
        total += count
        ws.cell(row, 3).value = count
        ws.cell(row, 4).value = count
        ws.cell(row, 5).value = "Đã bổ sung TC UAT/verify bug và task đối tác; đợt cuối PASS trong phạm vi sau loại hạ tầng."
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 2).value).strip().lower() == "17 module":
            ws.cell(row, 3).value = total
            ws.cell(row, 4).value = total
            ws.cell(row, 5).value = "Tổng raw sau khi bổ sung TC UAT/verify bug/task đối tác."


def main() -> None:
    wb = load_workbook(XLSX)
    covered = existing_partner_stt(wb)
    added = []
    for item in parse_rows():
        stt = int(item["stt"])
        if stt in covered:
            continue
        sheet = MODULE_MAP.get(str(item.get("Nhóm chức năng", "")).strip())
        if not sheet:
            raise RuntimeError(f"Missing module mapping for STT {stt}: {item.get('Nhóm chức năng')}")
        ws = wb[sheet]
        target_row = ws.max_row + 1
        source_row = max(4, target_row - 1)
        note = str(item.get("Mô tả", "")).strip()
        ws.append(
            [
                "Task/Bug đối tác",
                f"DTAC-STT-{stt:03d}",
                f"{item.get('Phân loại')} | {item.get('Độ ưu tiên')} | {item.get('Trạng thái xử lý')}",
                note.splitlines()[0][:250] if note else f"STT {stt}",
                "Có tài khoản/role và dữ liệu test theo danh sách Bug/Task đối tác.",
                f"STT {stt} - {item.get('Nhóm chức năng')}",
                "Thực hiện theo mô tả/mong muốn của đối tác; kiểm tra UI/API/export nếu áp dụng.",
                note,
            ]
        )
        copy_row_style(ws, source_row, target_row)
        covered.add(stt)
        added.append((sheet, f"DTAC-STT-{stt:03d}"))
    update_summary(wb)
    wb.save(XLSX)
    print({"added": len(added), "items": added})


if __name__ == "__main__":
    main()
