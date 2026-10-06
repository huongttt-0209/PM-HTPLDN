from __future__ import annotations

import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment


ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = ROOT / "output" / "test-cases"
XLSX = ROOT / "output" / "bao-cao-tong-hop-qa" / "Test-case.xlsx"

FOLDER_TO_SHEET = {
    "CG-TVV": "05. Chuyên gia tư vấn",
    "QTHT": "13. Quản trị hệ thống",
    "hoi-dap": "03. Hỏi đáp pháp lý",
    "doanh-nghiep": "10. Quản lý doanh nghiệp",
}

SUMMARY_TO_SHEET = {
    "Dashboard tổng quan": "02. Dashboard tổng quan",
    "Hỏi đáp pháp lý": "03. Hỏi đáp pháp lý",
    "Đào tạo tập huấn": "04. Đào tạo tập huấn",
    "Chuyên gia tư vấn": "05. Chuyên gia tư vấn",
    "Người hỗ trợ pháp luật": "06. Người hỗ trợ pháp luật",
    "Tổ chức tư vấn": "07. Tổ chức tư vấn",
    "Vụ việc HTPL": "08. Vụ việc HTPL",
    "Chi trả chi phí": "09. Chi trả chi phí",
    "Quản lý doanh nghiệp": "10. Quản lý doanh nghiệp",
    "Đánh giá": "11. Đánh giá",
    "Biểu mẫu": "12. Biểu mẫu",
    "Quản trị hệ thống": "13. Quản trị hệ thống",
    "Báo cáo thống kê": "14. Báo cáo thống kê",
    "Tư vấn chuyên sâu": "15. Tư vấn chuyên sâu",
    "Tư vấn nhanh": "16. Tư vấn nhanh",
    "Hợp đồng tư vấn": "17. Hợp đồng tư vấn",
    "Chương trình HTPLDN": "18. Chương trình HTPLDN",
}

ID_RE = re.compile(r"^\s*\|\s*`?(TC-[A-Za-z0-9_.-]+)`?\s*\|")


def is_execution_file(path: Path) -> bool:
    name = path.name.lower()
    if name.startswith("00-") or "review" in name or name.startswith("result-"):
        return False
    return "-tc-" in name or re.match(r"\d+-tc", name) is not None


def split_markdown_row(line: str) -> list[str]:
    return [cell.strip().strip("`") for cell in line.strip().strip("|").split("|")]


def parse_source_rows() -> dict[str, dict[str, str]]:
    rows = {}
    for path in sorted(SOURCE_ROOT.rglob("*.md")):
        if not is_execution_file(path):
            continue
        folder = path.relative_to(SOURCE_ROOT).parts[0]
        sheet = FOLDER_TO_SHEET.get(folder)
        if not sheet:
            continue
        for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
            if re.match(r"^\s*\|\s*-+", line):
                continue
            match = ID_RE.match(line)
            if not match:
                continue
            cells = split_markdown_row(line)
            if len(cells) < 7:
                continue
            tc_id = match.group(1).strip()
            rows[tc_id] = {
                "sheet": sheet,
                "folder": folder,
                "file": str(path.relative_to(ROOT)),
                "trace": cells[1] if len(cells) > 1 else "",
                "title": cells[2] if len(cells) > 2 else "",
                "precondition": cells[3] if len(cells) > 3 else "",
                "data": cells[4] if len(cells) > 4 else "",
                "steps": cells[5] if len(cells) > 5 else "",
                "expected": cells[6] if len(cells) > 6 else "",
                "type": cells[7] if len(cells) > 7 else "",
            }
    return rows


def existing_ids(wb) -> set[str]:
    ids = set()
    for ws in wb.worksheets:
        if ws.title == "Tổng hợp":
            continue
        for values in ws.iter_rows(min_row=4, min_col=2, max_col=2, values_only=True):
            if values[0]:
                ids.add(str(values[0]).strip())
    return ids


def copy_row_style(ws, source_row: int, target_row: int) -> None:
    for col in range(1, 9):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = Alignment(wrap_text=True, vertical="top")
        dst.number_format = src.number_format


def update_summary(wb) -> None:
    ws = wb["Tổng hợp"]
    total = 0
    for row in range(2, ws.max_row + 1):
        module = ws.cell(row, 2).value
        if not module or str(module).strip().lower() == "17 module":
            continue
        sheet = SUMMARY_TO_SHEET.get(str(module))
        if not sheet:
            continue
        count = sum(1 for values in wb[sheet].iter_rows(min_row=4, min_col=2, max_col=2, values_only=True) if values[0])
        total += count
        ws.cell(row, 3).value = count
        ws.cell(row, 4).value = count
        ws.cell(row, 5).value = "Đã tổng hợp testcase theo module."
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 2).value).strip().lower() == "17 module":
            ws.cell(row, 3).value = total
            ws.cell(row, 4).value = total
            ws.cell(row, 5).value = "Tổng testcase sau khi bổ sung."


def main() -> None:
    source = parse_source_rows()
    wb = load_workbook(XLSX)
    ids = existing_ids(wb)
    added = []
    for tc_id, item in source.items():
        if tc_id in ids:
            continue
        ws = wb[item["sheet"]]
        target_row = ws.max_row + 1
        source_row = max(4, target_row - 1)
        ws.append(
            [
                f"Output test-cases/{item['folder']}",
                tc_id,
                item["trace"],
                item["title"],
                item["precondition"],
                item["data"],
                item["steps"],
                item["expected"],
            ]
        )
        copy_row_style(ws, source_row, target_row)
        ids.add(tc_id)
        added.append((item["sheet"], tc_id, item["file"]))
    update_summary(wb)
    wb.save(XLSX)
    by_sheet = {}
    for sheet, *_ in added:
        by_sheet[sheet] = by_sheet.get(sheet, 0) + 1
    print({"source_tc": len(source), "added": len(added), "by_sheet": by_sheet})


if __name__ == "__main__":
    main()
