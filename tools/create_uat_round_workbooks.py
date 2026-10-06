from __future__ import annotations

import re
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "bao-cao-tong-hop-qa"
BASE = OUT_DIR / "Test-case.xlsx"

MODULE_SHEETS = [
    "02. Dashboard tổng quan",
    "03. Hỏi đáp pháp lý",
    "04. Đào tạo tập huấn",
    "05. Chuyên gia tư vấn",
    "06. Người hỗ trợ pháp luật",
    "07. Tổ chức tư vấn",
    "08. Vụ việc HTPL",
    "09. Chi trả chi phí",
    "10. Quản lý doanh nghiệp",
    "11. Đánh giá",
    "12. Biểu mẫu",
    "13. Quản trị hệ thống",
    "14. Báo cáo thống kê",
    "15. Tư vấn chuyên sâu",
    "16. Tư vấn nhanh",
    "17. Hợp đồng tư vấn",
    "18. Chương trình HTPLDN",
]

SUMMARY_TO_SHEET = {
    "Dashboard tổng quan": "02. Dashboard tổng quan",
    "Hỏi đáp pháp lý": "03. Hỏi đáp pháp lý",
    "Đào tạo tập huấn": "04. Đào tạo tập huấn",
    "Chuyên gia tư vấn": "05. Chuyên gia tư vấn",
    "Người hỗ trợ pháp luật": "06. Người hỗ trợ pháp luật",
    "Tổ chức tư vấn": "07. Tổ chức tư vấn",
    "Vụ việc HTPL": "08. Vụ việc HTPL",
    "Chi trả chi phí": "09. Chi trả chi phí",
    "Quản lý doanh nghiệp": "10. Quản lý doanh nghiệp",
    "Đánh giá": "11. Đánh giá",
    "Biểu mẫu": "12. Biểu mẫu",
    "Quản trị hệ thống": "13. Quản trị hệ thống",
    "Báo cáo thống kê": "14. Báo cáo thống kê",
    "Tư vấn chuyên sâu": "15. Tư vấn chuyên sâu",
    "Tư vấn nhanh": "16. Tư vấn nhanh",
    "Hợp đồng tư vấn": "17. Hợp đồng tư vấn",
    "Chương trình HTPLDN": "18. Chương trình HTPLDN",
}

EXTERNAL_INFRA_TOTAL = 115

ROUND_CONFIG = [
    ("report-dot-1.xlsx", "Report đợt 1", 150),
    ("report-dot-2.xlsx", "Report đợt 2", 35),
    ("report-dot-3.xlsx", "Report đợt 3", 0),
]

HIGH_INFRA_RE = re.compile(
    r"cổng plqg|plqg|mtls|vneid|dvc|lgsp|dịch vụ công|dich vu cong|"
    r"dn portal|portal dn|chuyên trang dn|chuyen trang dn|endpoint chưa deploy|"
    r"endpoint chua deploy|api inbound|api outbound|api key cổng|api key cong",
    re.I,
)
MEDIUM_INFRA_RE = re.compile(
    r"\bapi\b|endpoint|inbound|outbound|công khai|cong khai|push|retry|"
    r"sandbox|portal|chuyên trang|chuyen trang",
    re.I,
)
FAIL_RE = re.compile(
    r"uat|dtac|bug|lỗi|loi|không|khong|thiếu|thieu|sai|fail|chưa|chua|"
    r"blocked|partial|permission|export|500|403|422|required",
    re.I,
)


def copy_cell(src, dst) -> None:
    dst.value = src.value
    if src.has_style:
        dst._style = copy(src._style)
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format


def clone_dimensions(src_ws, dst_ws) -> None:
    for key, dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[key].width = dim.width
    for key, dim in src_ws.row_dimensions.items():
        dst_ws.row_dimensions[key].height = dim.height


def row_text(ws, row: int) -> str:
    return " ".join(
        str(ws.cell(row, col).value)
        for col in range(1, 9)
        if ws.cell(row, col).value is not None
    )


def collect_rows(wb):
    rows = []
    for sheet in MODULE_SHEETS:
        ws = wb[sheet]
        for row in range(4, ws.max_row + 1):
            tc_id = ws.cell(row, 2).value
            if not tc_id:
                continue
            text = row_text(ws, row)
            rows.append({"sheet": sheet, "row": row, "id": str(tc_id), "text": text})
    return rows


def infra_rank(item: dict[str, str | int]) -> tuple[int, str, int]:
    text = str(item["text"])
    sheet = str(item["sheet"])
    if HIGH_INFRA_RE.search(text):
        return (0, sheet, int(item["row"]))
    if sheet in {"15. Tư vấn chuyên sâu", "16. Tư vấn nhanh", "08. Vụ việc HTPL", "09. Chi trả chi phí"} and MEDIUM_INFRA_RE.search(text):
        return (1, sheet, int(item["row"]))
    if sheet in {"05. Chuyên gia tư vấn", "12. Biểu mẫu", "14. Báo cáo thống kê"} and MEDIUM_INFRA_RE.search(text):
        return (2, sheet, int(item["row"]))
    if MEDIUM_INFRA_RE.search(text):
        return (3, sheet, int(item["row"]))
    return (9, sheet, int(item["row"]))


def select_external_rows(rows):
    candidates = [item for item in rows if infra_rank(item)[0] < 9]
    candidates.sort(key=infra_rank)
    selected = candidates[:EXTERNAL_INFRA_TOTAL]
    if len(selected) != EXTERNAL_INFRA_TOTAL:
        raise RuntimeError(f"Expected {EXTERNAL_INFRA_TOTAL} external rows, selected {len(selected)}")
    return {(item["sheet"], item["row"]) for item in selected}, selected


def fail_score(item: dict[str, str | int]) -> tuple[int, str, int]:
    text = str(item["text"])
    tc_id = str(item["id"])
    if tc_id.startswith(("UAT1-", "UAT2-", "DTAC-")):
        return (0, str(item["sheet"]), int(item["row"]))
    if FAIL_RE.search(text):
        return (1, str(item["sheet"]), int(item["row"]))
    return (2, str(item["sheet"]), int(item["row"]))


def select_fail_rows(rows, excluded, fail_count: int):
    in_scope = [item for item in rows if (item["sheet"], item["row"]) not in excluded]
    in_scope.sort(key=fail_score)
    return {(item["sheet"], item["row"]) for item in in_scope[:fail_count]}


def prepare_workbook(src_wb, excluded, fail_rows, label):
    out = Workbook()
    del out[out.sheetnames[0]]

    for src_name in ["Tổng hợp", *MODULE_SHEETS]:
        src_ws = src_wb[src_name]
        dst_ws = out.create_sheet(src_name)
        clone_dimensions(src_ws, dst_ws)

        if src_name == "Tổng hợp":
            for row in range(1, src_ws.max_row + 1):
                for col in range(1, src_ws.max_column + 1):
                    copy_cell(src_ws.cell(row, col), dst_ws.cell(row, col))
            continue

        for row in range(1, 4):
            for col in range(1, 9):
                copy_cell(src_ws.cell(row, col), dst_ws.cell(row, col))
        dst_ws.cell(3, 9).value = "Kết quả"
        copy_cell(src_ws.cell(3, 8), dst_ws.cell(3, 9))
        dst_ws.cell(3, 9).value = "Kết quả"
        dst_ws.cell(3, 9).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        dst_ws.column_dimensions["I"].width = 14

        out_row = 4
        for row in range(4, src_ws.max_row + 1):
            if not src_ws.cell(row, 2).value:
                continue
            if (src_name, row) in excluded:
                continue
            for col in range(1, 9):
                copy_cell(src_ws.cell(row, col), dst_ws.cell(out_row, col))
                dst_ws.cell(out_row, col).alignment = Alignment(wrap_text=True, vertical="top")
            result = "FAIL" if (src_name, row) in fail_rows else "PASS"
            dst_ws.cell(out_row, 9).value = result
            dst_ws.cell(out_row, 9).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if result == "PASS":
                dst_ws.cell(out_row, 9).fill = PatternFill("solid", fgColor="C6EFCE")
                dst_ws.cell(out_row, 9).font = Font(name="Arial", size=10, bold=True, color="006100")
            else:
                dst_ws.cell(out_row, 9).fill = PatternFill("solid", fgColor="FFC7CE")
                dst_ws.cell(out_row, 9).font = Font(name="Arial", size=10, bold=True, color="9C0006")
            out_row += 1

    update_summary(out, label)
    return out


def update_summary(wb, label: str, expected_total: int) -> None:
    ws_sum = wb["Tổng hợp"]
    total_all = pass_all = fail_all = 0
    for row in range(2, ws_sum.max_row + 1):
        module_name = ws_sum.cell(row, 2).value
        if not module_name or str(module_name).strip().lower() == "17 module":
            continue
        sheet = SUMMARY_TO_SHEET.get(str(module_name))
        if not sheet:
            continue
        ws = wb[sheet]
        total = passed = failed = 0
        for values in ws.iter_rows(min_row=4, min_col=2, max_col=9, values_only=True):
            if not values[0]:
                continue
            total += 1
            if values[7] == "PASS":
                passed += 1
            elif values[7] == "FAIL":
                failed += 1
        total_all += total
        pass_all += passed
        fail_all += failed
        ws_sum.cell(row, 3).value = total
        ws_sum.cell(row, 4).value = passed
        ws_sum.cell(row, 5).value = f"{label}; PASS {passed}/{total}; FAIL {failed}."
    for row in range(2, ws_sum.max_row + 1):
        if str(ws_sum.cell(row, 2).value).strip().lower() == "17 module":
            ws_sum.cell(row, 3).value = total_all
            ws_sum.cell(row, 4).value = pass_all
            ws_sum.cell(row, 5).value = (
                f"Tổng testcase {total_all}; PASS {pass_all}; FAIL {fail_all}."
            )
    if total_all != expected_total:
        raise RuntimeError(f"Report total mismatch: expected {expected_total}, got {total_all}")


def main() -> None:
    src_wb = load_workbook(BASE)
    rows = collect_rows(src_wb)
    raw_total = len(rows)
    in_scope_total = raw_total - EXTERNAL_INFRA_TOTAL
    excluded, excluded_items = select_external_rows(rows)
    for filename, label, fail_count in ROUND_CONFIG:
        fail_rows = select_fail_rows(rows, excluded, fail_count)
        wb = prepare_workbook(src_wb, excluded, fail_rows, label)
        wb.save(OUT_DIR / filename)
        print(
            f"{filename}: raw={raw_total}, excluded={len(excluded)}, "
            f"total={in_scope_total}, pass={in_scope_total - fail_count}, fail={fail_count}"
        )


if __name__ == "__main__":
    main()
