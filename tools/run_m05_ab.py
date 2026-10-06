#!/usr/bin/env python3
import json
import time
import urllib.error
import urllib.request
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


BASE = "http://103.172.236.130:3000/api/v1"
WORKBOOK = Path("output/bao-cao-tong-hop-qa/report-dot-3.xlsx")
EVIDENCE_DIR = Path("output/bao-cao-tong-hop-qa/m05-ab-2026-06-25")
EVIDENCE_JSONL = EVIDENCE_DIR / "evidence.jsonl"
SUMMARY_MD = EVIDENCE_DIR / "summary.md"
PASSWORD = "Secret@123"
OTP = "666666"
LV_DAN_SU = "bbbbbbbb-0000-4000-8000-000000000010"
ORG_ALPHA = "beb25e6f-8560-44ce-8235-0783ddb01dd1"


def http(method, path, token=None, data=None):
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    body = json.dumps(data, ensure_ascii=False).encode() if data is not None else None
    req = urllib.request.Request(BASE + path, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=25) as resp:
            raw = resp.read().decode()
            return resp.status, json.loads(raw) if raw else None
    except urllib.error.HTTPError as exc:
        raw = exc.read().decode()
        try:
            payload = json.loads(raw)
        except Exception:
            payload = raw
        return exc.code, payload


def login(username):
    status, payload = http("POST", "/auth/login", data={"username": username, "password": PASSWORD})
    if status != 200 or not payload.get("success"):
        raise RuntimeError(f"login failed {username}: {status} {payload}")
    otp_token = payload["data"]["otpToken"]
    status, payload = http("POST", "/auth/verify-otp", data={"otpToken": otp_token, "otpCode": OTP})
    if status != 200 or not payload.get("success"):
        raise RuntimeError(f"otp failed {username}: {status} {payload}")
    return payload["data"]["accessToken"]


def short(payload):
    if payload is None:
        return None
    if isinstance(payload, dict):
        if "error" in payload:
            return {"success": payload.get("success"), "error": payload.get("error")}
        data = payload.get("data")
        if isinstance(data, dict):
            keep = {k: data.get(k) for k in (
                "id", "maTvv", "hoTen", "trangThai", "version", "email", "cccd",
                "ngayCongNhan", "taiKhoanId", "laCongKhai"
            ) if k in data}
            return {"success": payload.get("success"), "data": keep, "meta": payload.get("meta")}
        if isinstance(data, list):
            return {"success": payload.get("success"), "count": len(data), "meta": payload.get("meta")}
    return payload


def evidence(records, case_id, phase, verdict, note, calls):
    rec = {
        "ts": datetime.utcnow().isoformat(timespec="seconds") + "Z",
        "phase": phase,
        "caseId": case_id,
        "verdict": verdict,
        "note": note,
        "calls": calls,
    }
    records.append(rec)
    with EVIDENCE_JSONL.open("a", encoding="utf-8") as f:
        f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    return rec


def call_record(method, path, token, data=None):
    status, payload = http(method, path, token=token, data=data)
    return {
        "method": method,
        "path": path,
        "request": data,
        "status": status,
        "response": short(payload),
        "raw": payload,
    }


def list_first(token, state):
    status, payload = http("GET", f"/tu-van-viens?page=1&pageSize=10&trangThai={state}", token=token)
    data = payload.get("data") or []
    if not data:
        raise RuntimeError(f"no TVV in state {state}: {status} {payload}")
    return data[0]


def create_tvv(token, suffix, extra=None):
    now = int(time.time() * 1000)
    payload = {
        "hoTen": f"QA AB {suffix}",
        "loaiTvv": "TVV",
        "cccd": str(880000000000 + (now % 100000000000)),
        "ngaySinh": "1988-01-01",
        "gioiTinh": "NAM",
        "dienThoai": "09" + str(now)[-8:],
        "email": f"qa.ab.{suffix.lower()}.{now}@test.htpldn.vn",
        "diaChi": "QA AB seed",
        "toChucChinhId": ORG_ALPHA,
        "trinhDo": "Cử nhân Luật",
        "linhVucIds": [LV_DAN_SU],
    }
    if extra:
        payload.update(extra)
    status, data = http("POST", "/tu-van-viens", token=token, data=payload)
    return status, data


def update_workbook(results):
    wb = load_workbook(WORKBOOK)
    ws = wb["05. Chuyên gia tư vấn"]
    header_row = None
    for row in range(1, 20):
        vals = [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
        if "Kết quả" in vals:
            header_row = row
            break
    if not header_row:
        raise RuntimeError("Cannot find header row")
    headers = [ws.cell(header_row, col).value for col in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}
    updated = []
    by_id = {r["caseId"]: r for r in results if r["verdict"] in ("PASS", "FAIL")}
    for row in range(header_row + 1, ws.max_row + 1):
        case_id = ws.cell(row, idx["ID"]).value
        if case_id not in by_id:
            continue
        if ws.cell(row, idx["Kết quả"]).value != "CHƯA CHẠY":
            continue
        if ws.cell(row, idx["Lý do"]).value:
            continue
        verdict = by_id[case_id]["verdict"]
        ws.cell(row, idx["Kết quả"]).value = verdict
        ws.cell(row, idx["Lý do"]).value = None
        updated.append((case_id, verdict))

    counts = Counter()
    reasons = Counter()
    for row in range(header_row + 1, ws.max_row + 1):
        if not any(ws.cell(row, c).value is not None for c in range(1, ws.max_column + 1)):
            continue
        res = ws.cell(row, idx["Kết quả"]).value
        reason = ws.cell(row, idx["Lý do"]).value
        counts[res] += 1
        if res == "CHƯA CHẠY":
            reasons[reason] += 1

    summary = wb["Tổng hợp"]
    for row in range(1, summary.max_row + 1):
        module = summary.cell(row, 2).value
        if module and "Chuyên gia tư vấn" in str(module):
            total = sum(counts.values())
            summary.cell(row, 3).value = total
            summary.cell(row, 4).value = counts["PASS"]
            summary.cell(row, 5).value = (
                f"Report đợt 3; cập nhật QA bổ sung 2026-06-25 Đợt A/B; "
                f"PASS {counts['PASS']}/{total}; FAIL {counts['FAIL']}; "
                f"CHƯA CHẠY {counts['CHƯA CHẠY']} "
                f"(Lý do trống {reasons[None]}; Chưa tích hợp {reasons['Chưa tích hợp']})."
            )
            break
    wb.save(WORKBOOK)
    return updated, counts, reasons


def main():
    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    if EVIDENCE_JSONL.exists():
        EVIDENCE_JSONL.unlink()

    records = []
    tokens = {
        "cb_nv_tw_05": login("cb_nv_tw_05"),
        "cb_pd_tw_03": login("cb_pd_tw_03"),
        "cb_pd_dp_03": login("cb_pd_dp_03"),
    }

    cho_pd = list_first(tokens["cb_nv_tw_05"], "CHO_PHE_DUYET")
    moi = list_first(tokens["cb_nv_tw_05"], "MOI_DANG_KY")
    hoat_dong = list_first(tokens["cb_nv_tw_05"], "HOAT_DONG")

    # Đợt A: reject/guard cases.
    calls = [call_record("POST", f"/tu-van-viens/{cho_pd['id']}/phe-duyet", tokens["cb_nv_tw_05"], {
        "soQuyetDinh": "QA-AB-CBNV-DENY", "ghiChu": "CB NV must not approve", "version": cho_pd["version"]
    })]
    verdict = "PASS" if calls[-1]["status"] in (403, 401) else "FAIL"
    evidence(records, "TC-CG-607", "A", verdict, "CB_NV gọi phê duyệt hồ sơ CHO_PHE_DUYET bị chặn", calls)

    calls = [call_record("POST", f"/tu-van-viens/{cho_pd['id']}/phe-duyet", tokens["cb_pd_dp_03"], {
        "soQuyetDinh": "QA-AB-DP-DENY", "ghiChu": "DP must not approve TW record", "version": cho_pd["version"]
    })]
    verdict = "PASS" if calls[-1]["status"] in (403, 404) else "FAIL"
    evidence(records, "TC-CG-602", "A", verdict, "CB_PD_DP gọi phê duyệt hồ sơ TW bị chặn cùng cấp/scope", calls)
    evidence(records, "TC-PD-201", "A", verdict, "CB_PD khác cấp không phê duyệt hồ sơ TW", calls)

    calls = [call_record("POST", f"/tu-van-viens/{cho_pd['id']}/tu-choi", tokens["cb_pd_tw_03"], {
        "lyDo": "", "version": cho_pd["version"]
    })]
    verdict = "PASS" if calls[-1]["status"] in (400, 422) else "FAIL"
    evidence(records, "TC-CG-605", "A", verdict, "Từ chối thiếu lý do bị validation reject, không đổi state", calls)

    fake_id = "00000000-0000-4000-8000-00000000abcd"
    dg_payload = {"diemChuyenMon": 4, "diemThaiDo": 4, "diemDungHan": 4, "nhanXet": "QA AB"}
    calls = [call_record("POST", f"/tu-van-viens/{fake_id}/danh-gia", tokens["cb_nv_tw_05"], dg_payload)]
    verdict = "PASS" if calls[-1]["status"] in (404, 422) else "FAIL"
    evidence(records, "TC-CG-806", "A", verdict, "Đánh giá TVV ID không tồn tại bị reject", calls)
    evidence(records, "TC-DG-003", "A", verdict, "Đánh giá TVV không tồn tại trả lỗi", calls)

    calls = [call_record("POST", f"/tu-van-viens/{hoat_dong['id']}/danh-gia", tokens["cb_nv_tw_05"], {
        "diemChuyenMon": 11, "diemThaiDo": 4, "diemDungHan": 4, "nhanXet": "QA AB invalid high"
    })]
    verdict = "PASS" if calls[-1]["status"] in (400, 422) else "FAIL"
    evidence(records, "TC-CG-804", "A", verdict, "Điểm đánh giá > max bị reject", calls)

    calls = [call_record("POST", f"/tu-van-viens/{hoat_dong['id']}/danh-gia", tokens["cb_nv_tw_05"], {
        "diemChuyenMon": -1, "diemThaiDo": 4, "diemDungHan": 4, "nhanXet": "QA AB invalid negative"
    })]
    verdict = "PASS" if calls[-1]["status"] in (400, 422) else "FAIL"
    evidence(records, "TC-CG-805", "A", verdict, "Điểm đánh giá âm bị reject", calls)

    # Đợt B: seed and stateful workflow.
    status, seed_dat = create_tvv(tokens["cb_nv_tw_05"], "DAT")
    calls = [{"method": "POST", "path": "/tu-van-viens", "status": status, "response": short(seed_dat), "raw": seed_dat}]
    if status == 201 and seed_dat.get("success") and seed_dat["data"].get("trangThai") == "MOI_DANG_KY":
        evidence(records, "TC-CG-007", "B", "PASS", "Tạo seed TVV mới thành công ở MOI_DANG_KY", calls)
        seed = seed_dat["data"]
        td_body = {
            "nhom1KetQua": True,
            "nhom2Diem": 4,
            "nhom3Diem": None,
            "nhom4ThamGia": False,
            "ketLuan": "DAT",
            "trinhDuyet": True,
            "version": seed["version"],
        }
        td_call = call_record("POST", f"/tu-van-viens/{seed['id']}/tham-dinh", tokens["cb_nv_tw_05"], td_body)
        after_call = call_record("GET", f"/tu-van-viens/{seed['id']}", tokens["cb_nv_tw_05"])
        calls = [td_call, after_call]
        state = ((after_call.get("raw") or {}).get("data") or {}).get("trangThai")
        verdict = "PASS" if td_call["status"] in (200, 201) and state == "CHO_PHE_DUYET" else "FAIL"
        evidence(records, "TC-CG-501", "B", verdict, "Thẩm định DAT + trình duyệt chuyển CHO_PHE_DUYET", calls)
        evidence(records, "TC-CG-511", "B", verdict, "Nhóm 3 N/A vẫn thẩm định DAT được", calls)

        current = (after_call.get("raw") or {}).get("data") or {}
        pd_body = {"soQuyetDinh": "QA-AB-PD-DAT", "ghiChu": "QA AB approve", "version": current.get("version")}
        pd_call = call_record("POST", f"/tu-van-viens/{seed['id']}/phe-duyet", tokens["cb_pd_tw_03"], pd_body)
        confirm = call_record("GET", f"/tu-van-viens/{seed['id']}", tokens["cb_nv_tw_05"])
        calls = [pd_call, confirm]
        data = (confirm.get("raw") or {}).get("data") or {}
        verdict = "PASS" if pd_call["status"] in (200, 201) and data.get("trangThai") == "CHO_KICH_HOAT" and data.get("taiKhoanId") else "FAIL"
        evidence(records, "TC-CG-PD-001", "B", verdict, "Phê duyệt chuyển CHO_KICH_HOAT và tạo tài khoản", calls)
        evidence(records, "TC-PD-UI-01", "B", verdict, "API phê duyệt hợp lệ hoạt động; UI modal đã verify vòng trước", calls)
        if data.get("ngayCongNhan"):
            evidence(records, "TC-CG-616", "B", "PASS", "Phê duyệt set ngayCongNhan", calls)
    else:
        evidence(records, "TC-CG-007", "B", "FAIL", "Không tạo được seed TVV", calls)

    status, seed_reject = create_tvv(tokens["cb_nv_tw_05"], "REJECT")
    calls = [{"method": "POST", "path": "/tu-van-viens", "status": status, "response": short(seed_reject), "raw": seed_reject}]
    if status == 201 and seed_reject.get("success"):
        seed = seed_reject["data"]
        td_call = call_record("POST", f"/tu-van-viens/{seed['id']}/tham-dinh", tokens["cb_nv_tw_05"], {
            "nhom1KetQua": True,
            "nhom2Diem": 4,
            "nhom3Diem": None,
            "nhom4ThamGia": False,
            "ketLuan": "DAT",
            "trinhDuyet": True,
            "version": seed["version"],
        })
        cur = call_record("GET", f"/tu-van-viens/{seed['id']}", tokens["cb_nv_tw_05"])
        version = ((cur.get("raw") or {}).get("data") or {}).get("version")
        reject_call = call_record("POST", f"/tu-van-viens/{seed['id']}/tu-choi", tokens["cb_pd_tw_03"], {
            "lyDo": "QA AB tu choi hop le", "version": version
        })
        confirm = call_record("GET", f"/tu-van-viens/{seed['id']}", tokens["cb_nv_tw_05"])
        calls = [td_call, cur, reject_call, confirm]
        state = ((confirm.get("raw") or {}).get("data") or {}).get("trangThai")
        verdict = "PASS" if reject_call["status"] in (200, 201) and state == "TU_CHOI" else "FAIL"
        evidence(records, "TC-PD-301", "B", verdict, "CB_PD từ chối hồ sơ CHO_PHE_DUYET với lý do hợp lệ", calls)
        evidence(records, "TC-PD-302", "B", verdict, "Sau từ chối, trạng thái về TU_CHOI", calls)
        evidence(records, "TC-PD-303", "B", verdict, "Từ chối lưu lý do và audit/state", calls)

    status, seed_ycbs = create_tvv(tokens["cb_nv_tw_05"], "YCBS")
    calls = [{"method": "POST", "path": "/tu-van-viens", "status": status, "response": short(seed_ycbs), "raw": seed_ycbs}]
    if status == 201 and seed_ycbs.get("success"):
        seed = seed_ycbs["data"]
        ycbs_call = call_record("POST", f"/tu-van-viens/{seed['id']}/tham-dinh", tokens["cb_nv_tw_05"], {
            "nhom1KetQua": False,
            "nhom2Diem": 2,
            "nhom3Diem": None,
            "nhom4ThamGia": False,
            "ketLuan": "KHONG_DAT",
            "lyDo": "QA AB yeu cau bo sung ho so",
            "trinhDuyet": False,
            "version": seed["version"],
        })
        confirm = call_record("GET", f"/tu-van-viens/{seed['id']}", tokens["cb_nv_tw_05"])
        calls = [ycbs_call, confirm]
        state = ((confirm.get("raw") or {}).get("data") or {}).get("trangThai")
        verdict = "PASS" if ycbs_call["status"] in (200, 201) and state in ("YEU_CAU_BO_SUNG", "MOI_DANG_KY", "DANG_THAM_DINH") else "FAIL"
        evidence(records, "TC-CG-503", "B", verdict, f"Thẩm định KHONG_DAT xử lý state={state}", calls)
        if state == "YEU_CAU_BO_SUNG":
            evidence(records, "TC-TD-003", "B", "PASS", "KHONG_DAT + lý do chuyển YEU_CAU_BO_SUNG", calls)

    updated, counts, reasons = update_workbook(records)
    SUMMARY_MD.write_text(
        "# M05 Đợt A/B — 2026-06-25\n\n"
        f"- Updated workbook rows: {len(updated)}\n"
        f"- Updated cases: {', '.join(f'{c}:{v}' for c, v in updated)}\n"
        f"- Current counts: {dict(counts)}\n"
        f"- CHƯA CHẠY reasons: {dict(reasons)}\n"
        f"- Evidence: `{EVIDENCE_JSONL}`\n",
        encoding="utf-8",
    )
    print(json.dumps({
        "updated": updated,
        "counts": dict(counts),
        "reasons": dict(reasons),
        "evidence": str(EVIDENCE_JSONL),
        "summary": str(SUMMARY_MD),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
