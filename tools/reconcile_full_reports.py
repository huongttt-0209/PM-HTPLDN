from __future__ import annotations

import re
from collections import Counter, defaultdict
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "bao-cao-tong-hop-qa"
MASTER = OUT / "Test-case.xlsx"
HUONG_TC = ROOT / "output" / "test-cases" / "huong" / "huong" / "test-cases"
PASS_SOURCE = ROOT / "output" / "test-cases" / "c-Hoa"

REPORTS = [
    ("report-dot-1.xlsx", "Report đợt 1"),
    ("report-dot-2.xlsx", "Report đợt 2"),
    ("report-dot-3.xlsx", "Report đợt 3"),
]

FOLDER_TO_SHEET = {
    "dashboard": "02. Dashboard tổng quan",
    "hoi-dap": "03. Hỏi đáp pháp lý",
    "dao-tao": "04. Đào tạo tập huấn",
    "CG-TVV": "05. Chuyên gia tư vấn",
    "vu-viec": "08. Vụ việc HTPL",
    "chi-tra": "09. Chi trả chi phí",
    "quan-ly-doanh-nghiep": "10. Quản lý doanh nghiệp",
    "danh-gia": "11. Đánh giá",
    "bieu-mau": "12. Biểu mẫu",
    "QTHT": "13. Quản trị hệ thống",
    "bao-cao-tk": "14. Báo cáo thống kê",
    "tv-chuyen-sau": "15. Tư vấn chuyên sâu",
    "tv-nhanh": "16. Tư vấn nhanh",
    "hop-dong-tv": "17. Hợp đồng tư vấn",
    "ct-htpldn-gd1": "18. Chương trình HTPLDN",
    "ct-htpldn-gd2": "18. Chương trình HTPLDN",
}

FOLDER_PREFIX = {
    "dashboard": "DASH",
    "hoi-dap": "HD",
    "dao-tao": "DT",
    "CG-TVV": "CG",
    "vu-viec": "VV",
    "chi-tra": "CT",
    "quan-ly-doanh-nghiep": "DN",
    "danh-gia": "DG",
    "bieu-mau": "BM",
    "QTHT": "QTHT",
    "bao-cao-tk": "BC",
    "tv-chuyen-sau": "TVCS",
    "tv-nhanh": "TVN",
    "hop-dong-tv": "HDTV",
    "ct-htpldn-gd1": "CTHT1",
    "ct-htpldn-gd2": "CTHT2",
}

SUMMARY_TO_SHEET = {
    "Dashboard tổng quan": "02. Dashboard tổng quan",
    "Hỏi đáp pháp lý": "03. Hỏi đáp pháp lý",
    "Đào tạo tập huấn": "04. Đào tạo tập huấn",
    "Chuyên gia tư vấn": "05. Chuyên gia tư vấn",
    "Người hỗ trợ pháp luật": "06. Người hỗ trợ pháp luật",
    "Tổ chức tư vấn": "07. Tổ chức tư vấn",
    "Vụ việc HTPL": "08. Vụ việc HTPL",
    "Chi trả chi phí": "09. Chi trả chi phí",
    "Quản lý doanh nghiệp": "10. Quản lý doanh nghiệp",
    "Đánh giá": "11. Đánh giá",
    "Biểu mẫu": "12. Biểu mẫu",
    "Quản trị hệ thống": "13. Quản trị hệ thống",
    "Báo cáo thống kê": "14. Báo cáo thống kê",
    "Tư vấn chuyên sâu": "15. Tư vấn chuyên sâu",
    "Tư vấn nhanh": "16. Tư vấn nhanh",
    "Hợp đồng tư vấn": "17. Hợp đồng tư vấn",
    "Chương trình HTPLDN": "18. Chương trình HTPLDN",
}

MODULE_SHEETS = list(SUMMARY_TO_SHEET.values())

ID_RE = re.compile(r"^\s*\|\s*`?([A-Z][A-Z0-9_.-]*-[A-Z0-9_.-]*\d[A-Z0-9_.-]*)`?\s*\|")
INFRA_RE = re.compile(
    r"cổng plqg|cong plqg|plqg|mtls|vneid|dvc|lgsp|dịch vụ công|dich vu cong|"
    r"dn portal|portal dn|chuyên trang dn|chuyen trang dn|api inbound|api outbound|"
    r"endpoint|sandbox|tích hợp|tich hop",
    re.I,
)


def is_tc_file(path: Path) -> bool:
    name = path.name.lower()
    if (
        name.startswith("00-")
        or "review" in name
        or "traceability" in name
        or "trace-matrix" in name
        or "quality" in name
        or "filter-log" in name
        or "codex" in name
        or "sync-log" in name
        or name.endswith("-report.md")
    ):
        return False
    return "-tc-" in name or re.match(r"\d+-tc", name) is not None


def split_row(line: str) -> list[str]:
    return [cell.strip().strip("`") for cell in line.strip().strip("|").split("|")]


def parse_huong_rows() -> list[dict[str, str]]:
    rows = []
    for path in sorted(HUONG_TC.rglob("*.md")):
        if not is_tc_file(path):
            continue
        rel = path.relative_to(HUONG_TC)
        folder = rel.parts[0]
        sheet = FOLDER_TO_SHEET.get(folder)
        if not sheet:
            continue
        for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
            if re.match(r"^\s*\|\s*-+", line):
                continue
            match = ID_RE.match(line)
            if not match:
                continue
            cells = split_row(line)
            if len(cells) < 7:
                continue
            source_id = match.group(1)
            if source_id in {"ID", "TC-ID"} or source_id.startswith("SPEC-CLARIFY"):
                continue
            rows.append(
                {
                    "source_id": source_id,
                    "folder": folder,
                    "sheet": sheet,
                    "file": str(rel),
                    "trace": cells[1] if len(cells) > 1 else "",
                    "title": cells[2] if len(cells) > 2 else "",
                    "precondition": cells[3] if len(cells) > 3 else "",
                    "data": cells[4] if len(cells) > 4 else "",
                    "steps": cells[5] if len(cells) > 5 else "",
                    "expected": cells[6] if len(cells) > 6 else "",
                    "type": cells[7] if len(cells) > 7 else "",
                }
            )
    return rows


def copy_cell(src, dst) -> None:
    dst.value = src.value
    if src.has_style:
        dst._style = copy(src._style)
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format


def copy_row_style(ws, source_row: int, target_row: int, cols: int = 8) -> None:
    for col in range(1, cols + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = Alignment(wrap_text=True, vertical="top")
        dst.number_format = src.number_format


def row_signature(item: dict[str, str]) -> tuple[str, str, str]:
    return (
        item["sheet"],
        " ".join(item["title"].lower().split()),
        " ".join(item["trace"].lower().split()),
    )


def workbook_ids_and_signatures(wb):
    ids = {}
    signatures = set()
    for ws in wb.worksheets:
        if ws.title == "Tổng hợp":
            continue
        for row in range(4, ws.max_row + 1):
            tc_id = ws.cell(row, 2).value
            if not tc_id:
                continue
            tc_id = str(tc_id).strip()
            ids[tc_id] = (ws.title, row)
            signatures.add(
                (
                    ws.title,
                    " ".join(str(ws.cell(row, 4).value or "").lower().split()),
                    " ".join(str(ws.cell(row, 3).value or "").lower().split()),
                )
            )
    return ids, signatures


def unique_id(item: dict[str, str], ids: dict[str, tuple[str, int]], source_seen: Counter) -> str:
    source_id = item["source_id"]
    if source_seen[source_id] == 1 and source_id not in ids:
        return source_id
    prefix = FOLDER_PREFIX[item["folder"]]
    if source_id.startswith("TC-"):
        candidate = f"TC-{prefix}-{source_id[3:]}"
    else:
        candidate = f"{prefix}-{source_id}"
    if candidate not in ids:
        return candidate
    n = 2
    while f"{candidate}-{n}" in ids:
        n += 1
    return f"{candidate}-{n}"


def choose_cluster(sheet: str, item: dict[str, str]) -> str:
    # Match the existing workbook style without creating new "Bổ sung" clusters.
    text = f"{item['source_id']} {item['trace']} {item['title']} {item['file']}".lower()
    fallback = {
        "02. Dashboard tổng quan": "Cluster 1 — Layout tổng quan & Header",
        "03. Hỏi đáp pháp lý": "Cluster 0 — Quản lý hỏi đáp (CRUD + Export + Batch + Sort/Fi",
        "04. Đào tạo tập huấn": "Cluster 0 — Kế hoạch đào tạo năm (KH năm)",
        "05. Chuyên gia tư vấn": "Cluster 1 — Quản lý danh sách & CRUD TVV (FR-IV-01)",
        "08. Vụ việc HTPL": "Cluster 0 — Quản lý vụ việc",
        "09. Chi trả chi phí": "Cluster 1 — Danh sách & tiếp nhận hồ sơ",
        "10. Quản lý doanh nghiệp": "Cluster B — CRUD Doanh nghiệp (FR-V.III-01)",
        "11. Đánh giá": "Cluster 0 — Quản lý đợt đánh giá",
        "12. Biểu mẫu": "Cluster 1 — Quản lý thư mục biểu mẫu",
        "13. Quản trị hệ thống": "Cluster 2 — Danh mục dùng chung",
        "14. Báo cáo thống kê": "Cluster A — UI / SCR-IX-01 verify",
        "15. Tư vấn chuyên sâu": "Cluster 0 — Quản lý danh sách & chi tiết TVCS",
        "16. Tư vấn nhanh": "Cluster 2 — Quản lý phiên tư vấn nhanh (CMS)",
        "17. Hợp đồng tư vấn": "Cluster 1 — Quản lý hợp đồng tư vấn",
        "18. Chương trình HTPLDN": "Cluster 1 — Giai đoạn 1: Quản lý CT HTPLDN (CRUD, Tìm kiếm,",
    }
    if sheet == "13. Quản trị hệ thống":
        if "cau-hinh" in text or "sla" in text or "mẫu phản hồi" in text or "ngày lễ" in text:
            return "Cluster 1 — Cấu hình hệ thống"
        if "nhat-ky" in text or "nhật ký" in text or "audit" in text:
            return "Cluster 3 — Nhật ký hệ thống"
        if "tai-khoan" in text or "tài khoản" in text or "vai-tro" in text or "permission" in text or "phân quyền" in text:
            return "Cluster 4 — Tài khoản & phân quyền"
        return "Cluster 2 — Danh mục dùng chung"
    if sheet == "03. Hỏi đáp pháp lý":
        if "tim-kiem" in text or "tìm kiếm" in text:
            return "Cluster 1 — Tìm kiếm tổng hợp (UC11/14/19)"
        if "tiep-nhan" in text or "tiếp nhận" in text:
            return "Cluster 2 — Tiếp nhận xử lý hỏi đáp (UC12)"
        if "phan-cong" in text or "phân công" in text:
            return "Cluster 4 — Phân công xử lý câu hỏi (UC15)"
        if "phan-hoi" in text or "phản hồi" in text:
            return "Cluster 5 — Phản hồi câu hỏi (UC16)"
        if "phe-duyet" in text or "duyệt" in text or "cong-khai" in text or "công khai" in text:
            return "Cluster 6 — Phê duyệt, công khai, đóng hồ sơ (UC17)"
    if sheet == "05. Chuyên gia tư vấn":
        if "tim-kiem" in text:
            return "Cluster 2 — Tìm kiếm & Export Excel (FR-IV-02)"
        if "dang-ky" in text:
            return "Cluster 3 — NHT đăng ký mạng lưới (FR-IV-03)"
        if "cap-nhat-nang-luc" in text:
            return "Cluster 4 — Cập nhật năng lực (FR-IV-04)"
        if "xem-chi-tiet" in text or "lich-su" in text:
            return "Cluster 5 — Xem chi tiết (FR-IV-05)"
        if "tham-dinh" in text:
            return "Cluster 6 — Thẩm định hồ sơ (FR-IV-06)"
        if "phe-duyet" in text:
            return "Cluster 7 — Phê duyệt TVV (FR-IV-07)"
        if "cong-khai" in text:
            return "Cluster 8 — Công khai mạng lưới (FR-IV-08)"
        if "danh-gia" in text:
            return "Cluster 9 — Đánh giá TVV (FR-IV-09 + CROSS-01)"
        if "permission" in text:
            return "Cluster 13 — Permission Matrix tổng hợp"
    if sheet == "04. Đào tạo tập huấn":
        if "ctdt" in text:
            return "Cluster 1 — Chương trình đào tạo (CTĐT)"
        if "khoa-hoc" in text:
            return "Cluster 2 — Khóa học quản lý (SM-KHOAHOC)"
        if "bai-giang" in text or "tai-lieu" in text:
            return "Cluster 3 — Bài giảng & Kho tài liệu"
        if "lich-hoc" in text:
            return "Cluster 4 — Lịch học (FR-III-22)"
        if "nhch" in text or "kiem-tra" in text:
            return "Cluster 5 — Ngân hàng câu hỏi & Đề kiểm tra"
        if "giang-vien" in text:
            return "Cluster 6 — Giảng viên"
        if "de-xuat" in text:
            return "Cluster 7 — Đề xuất đào tạo"
        if "dang-ky" in text:
            return "Cluster 8 — Đăng ký đào tạo"
        if "diem-danh" in text or "ket-qua" in text:
            return "Cluster 9 — Điểm danh & Kết quả học tập"
        if "permission" in text:
            return "Cluster 14 — Phân quyền & Multi-tenant"
    if sheet == "10. Quản lý doanh nghiệp":
        if "tim-kiem" in text:
            return "Cluster C — Tìm kiếm Doanh nghiệp (FR-V.III-02)"
        if "ho-so-phap-ly" in text or "hspl" in text:
            return "Cluster D — Tab Hồ sơ pháp lý DN (FR-X.1-04 / UC150)"
        if "lich-su" in text or "chi-tra" in text:
            return "Cluster E — Tab Lịch sử Hỗ trợ + Tab Hồ sơ Chi trả (read-onl"
        if "permission" in text:
            return "Cluster F — Permission Matrix (11 role × FR-07)"
    if sheet == "14. Báo cáo thống kê":
        if "export" in text or "xlsx" in text or "pdf" in text:
            return "Cluster G — Xuất file XLSX / PDF (TT 17/2025)"
        if "permission" in text:
            return "Cluster H — Phân quyền 2-tier (BR-AUTH-08 v3.5)"
        if "chart" in text or "bieu-do" in text:
            return "Cluster C — BC Đào tạo + CG/TVV + Đánh giá (UC129-UC133)"
    if sheet == "18. Chương trình HTPLDN" and ("gd2" in text or "bao-cao" in text or "báo cáo" in text):
        return "Cluster 2 — Giai đoạn 2: Đợt báo cáo CT HTPLDN (SM-DOT-BC, G"
    return fallback.get(sheet, "Cluster 0")


def parse_pass_ids() -> set[str]:
    ids = set()
    for path in PASS_SOURCE.glob("result-*-pass.md"):
        for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
            if re.match(r"^\s*\|\s*-+", line):
                continue
            match = ID_RE.match(line)
            if match and match.group(1) != "ID":
                ids.add(match.group(1))
    return ids


def is_infra_row(values: list[str]) -> bool:
    return bool(INFRA_RE.search(" ".join(values)))


def add_huong_to_master() -> dict[str, int]:
    wb = load_workbook(MASTER)
    ids, signatures = workbook_ids_and_signatures(wb)
    rows = parse_huong_rows()
    source_seen = Counter(item["source_id"] for item in rows)
    added_by_sheet = Counter()
    skipped_same = 0
    for item in rows:
        if row_signature(item) in signatures:
            skipped_same += 1
            continue
        tc_id = unique_id(item, ids, source_seen)
        sheet = item["sheet"]
        ws = wb[sheet]
        row = ws.max_row + 1
        source_trace = item["trace"]
        if tc_id != item["source_id"]:
            source_trace = f"{item['source_id']} | {source_trace}".strip(" |")
        ws.append(
            [
                choose_cluster(sheet, item),
                tc_id,
                source_trace,
                item["title"],
                item["precondition"],
                item["data"],
                item["steps"],
                item["expected"],
            ]
        )
        copy_row_style(ws, max(4, row - 1), row)
        ids[tc_id] = (sheet, row)
        signatures.add(row_signature(item))
        added_by_sheet[sheet] += 1
    update_master_summary(wb)
    wb.save(MASTER)
    return {"source_rows": len(rows), "added": sum(added_by_sheet.values()), "skipped_same": skipped_same, **dict(added_by_sheet)}


def update_master_summary(wb) -> None:
    ws = wb["Tổng hợp"]
    total = 0
    for row in range(2, ws.max_row + 1):
        module = ws.cell(row, 2).value
        if not module or str(module).strip().lower() == "17 module":
            continue
        sheet = SUMMARY_TO_SHEET.get(str(module))
        if not sheet:
            continue
        count = sum(1 for values in wb[sheet].iter_rows(min_row=4, min_col=2, max_col=2, values_only=True) if values[0])
        total += count
        ws.cell(row, 3).value = count
        ws.cell(row, 4).value = count
        ws.cell(row, 5).value = "Đã tổng hợp testcase chi tiết theo module."
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 2).value).strip().lower() == "17 module":
            ws.cell(row, 3).value = total
            ws.cell(row, 4).value = total
            ws.cell(row, 5).value = "Tổng testcase chi tiết."


def set_status_style(cell, status: str) -> None:
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if status == "PASS":
        cell.fill = PatternFill("solid", fgColor="C6EFCE")
        cell.font = Font(name="Arial", size=10, bold=True, color="006100")
    elif status == "FAIL":
        cell.fill = PatternFill("solid", fgColor="FFC7CE")
        cell.font = Font(name="Arial", size=10, bold=True, color="9C0006")
    else:
        cell.fill = PatternFill("solid", fgColor="E7E6E6")
        cell.font = Font(name="Arial", size=10, bold=True, color="666666")


def ensure_report_columns(ws) -> None:
    if ws.cell(3, 9).value != "Kết quả":
        copy_cell(ws.cell(3, 8), ws.cell(3, 9))
        ws.cell(3, 9).value = "Kết quả"
    if ws.cell(3, 10).value != "Lý do":
        copy_cell(ws.cell(3, 8), ws.cell(3, 10))
        ws.cell(3, 10).value = "Lý do"
    ws.column_dimensions["I"].width = 14
    ws.column_dimensions["J"].width = 28
    ws.cell(3, 9).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(3, 10).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def update_report_summary(wb, label: str) -> dict[str, int]:
    totals = Counter()
    ws_sum = wb["Tổng hợp"]
    for row in range(2, ws_sum.max_row + 1):
        module = ws_sum.cell(row, 2).value
        if not module or str(module).strip().lower() == "17 module":
            continue
        sheet = SUMMARY_TO_SHEET.get(str(module))
        if not sheet:
            continue
        ws = wb[sheet]
        counts = Counter()
        for r in range(4, ws.max_row + 1):
            if not ws.cell(r, 2).value:
                continue
            status = str(ws.cell(r, 9).value or "").strip()
            counts["total"] += 1
            if status == "PASS":
                counts["pass"] += 1
            elif status == "FAIL":
                counts["fail"] += 1
            elif status == "CHƯA CHẠY":
                counts["notrun"] += 1
            elif status in {"BLOCKED", "N/A"}:
                counts["blocked"] += 1
        totals.update(counts)
        ws_sum.cell(row, 3).value = counts["total"]
        ws_sum.cell(row, 4).value = counts["pass"]
        ws_sum.cell(row, 5).value = (
            f"{label}; PASS {counts['pass']}/{counts['total']}; "
            f"FAIL {counts['fail']}; CHƯA CHẠY {counts['notrun']}; BLOCKED/N/A {counts['blocked']}."
        )
    for row in range(2, ws_sum.max_row + 1):
        if str(ws_sum.cell(row, 2).value).strip().lower() == "17 module":
            ws_sum.cell(row, 3).value = totals["total"]
            ws_sum.cell(row, 4).value = totals["pass"]
            ws_sum.cell(row, 5).value = (
                f"Tổng testcase {totals['total']}; PASS {totals['pass']}; "
                f"FAIL {totals['fail']}; CHƯA CHẠY {totals['notrun']}; BLOCKED/N/A {totals['blocked']}."
            )
    return dict(totals)


def sync_reports() -> dict[str, dict[str, int]]:
    master = load_workbook(MASTER, data_only=False)
    pass_ids = parse_pass_ids()
    output = {}
    for report, label in REPORTS:
        wb = load_workbook(OUT / report)
        added = 0
        evidence_updated = 0
        infra_reason = 0
        for sheet in MODULE_SHEETS:
            if sheet not in wb.sheetnames or sheet not in master.sheetnames:
                continue
            ws = wb[sheet]
            mws = master[sheet]
            ensure_report_columns(ws)
            existing = {str(ws.cell(r, 2).value).strip(): r for r in range(4, ws.max_row + 1) if ws.cell(r, 2).value}
            for row in range(4, mws.max_row + 1):
                tc_id = mws.cell(row, 2).value
                if not tc_id:
                    continue
                tc_id = str(tc_id).strip()
                if tc_id not in existing:
                    target = ws.max_row + 1
                    for col in range(1, 9):
                        copy_cell(mws.cell(row, col), ws.cell(target, col))
                        ws.cell(target, col).alignment = Alignment(wrap_text=True, vertical="top")
                    status = "PASS" if tc_id in pass_ids else "CHƯA CHẠY"
                    reason = "Đã có evidence PASS" if status == "PASS" else "Chưa có evidence chạy"
                    row_values = [str(mws.cell(row, col).value or "") for col in range(1, 9)]
                    if status == "CHƯA CHẠY" and is_infra_row(row_values):
                        reason = "Chưa tích hợp"
                    ws.cell(target, 9).value = status
                    ws.cell(target, 10).value = reason
                    set_status_style(ws.cell(target, 9), status)
                    ws.cell(target, 10).alignment = Alignment(wrap_text=True, vertical="top")
                    existing[tc_id] = target
                    added += 1
                else:
                    target = existing[tc_id]
                    current = str(ws.cell(target, 9).value or "").strip()
                    row_values = [str(ws.cell(target, col).value or "") for col in range(1, 9)]
                    if current in {"", "CHƯA CHẠY"} and tc_id in pass_ids:
                        ws.cell(target, 9).value = "PASS"
                        ws.cell(target, 10).value = "Đã có evidence PASS"
                        set_status_style(ws.cell(target, 9), "PASS")
                        evidence_updated += 1
                    elif current == "CHƯA CHẠY" and is_infra_row(row_values):
                        ws.cell(target, 10).value = "Chưa tích hợp"
                        infra_reason += 1
                    elif current == "CHƯA CHẠY" and not ws.cell(target, 10).value:
                        ws.cell(target, 10).value = "Chưa có evidence chạy"
            for row in range(4, ws.max_row + 1):
                if ws.cell(row, 2).value and ws.cell(row, 9).value == "CHƯA CHẠY" and not ws.cell(row, 10).value:
                    ws.cell(row, 10).value = "Chưa có evidence chạy"
        totals = update_report_summary(wb, label)
        wb.save(OUT / report)
        output[report] = {"added": added, "evidence_updated": evidence_updated, "infra_reason": infra_reason, **totals}
    return output


def main() -> None:
    master_result = add_huong_to_master()
    report_result = sync_reports()
    print({"master": master_result, "reports": report_result})


if __name__ == "__main__":
    main()
